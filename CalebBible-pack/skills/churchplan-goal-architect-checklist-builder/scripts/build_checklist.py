#!/usr/bin/env python3
"""
build_checklist.py — churchplan-goal-architect 3축 목표 점검표 렌더러.

goals.json → ① goal-checklist.md (3축 목표 요약표 + 유지/중단 + 경고)
             ② goal-checklist.xlsx (분기 점검 칸 포함, openpyxl 없으면 .csv 폴백)

렌더 전 핵심 게이트를 재강제한다(cys-insight 패턴 — garbage-in을 렌더 단계에서도 차단):
  필수 블록 · 정성↔정량 한 쌍 · 과욕 상한. 위반 시 렌더 거부(종료코드 1).
전체 정밀 검증은 fidelity-guard의 validate_goals.py가 맡는다.

사용:  python3 build_checklist.py <goals.json> [출력디렉토리]
종료코드: 0 = 성공, 1 = 게이트 위반(렌더 거부), 2 = 파일/JSON 오류
표준 라이브러리만 사용(xlsx는 openpyxl 있을 때만, 없으면 csv).
"""
import csv
import json
import os
import sys


def _d(x):
    return x if isinstance(x, dict) else {}


def _l(x):
    return x if isinstance(x, list) else []


def _ld(x):
    return [v for v in x if isinstance(v, dict)] if isinstance(x, list) else []


AXES = {"영성", "전도", "양육"}
GOALS_PER_AXIS_MAX = 2
GOALS_TOTAL_MAX = 6


def gate(g):
    """렌더 전 핵심 게이트 — fidelity-guard 본질 3종만 재강제."""
    errs = []
    for blk in ("meta", "inputs_echo", "axes"):
        if blk not in g:
            errs.append(f"필수 블록 '{blk}' 없음")
    # 최상위 dict/list-기대 필드가 비-dict/list로 오염되면 형식오류 등록(조용한 통과 금지·검증기 parity)
    for f in ("meta", "inputs_echo", "confirmed"):
        if f in g and g.get(f) is not None and not isinstance(g.get(f), dict):
            errs.append(f"최상위 '{f}'가 객체가 아님 — 형식오류(렌더 거부)")
    for f in ("axes", "keep_list", "stop_list", "warnings"):
        if f in g and not isinstance(g.get(f), list):
            errs.append(f"최상위 '{f}'가 배열이 아님 — 형식오류(렌더 거부)")
    ie_g = g.get("inputs_echo")
    if isinstance(ie_g, dict):
        for f in ("core_scriptures", "needs_care"):
            if f in ie_g and not isinstance(ie_g.get(f), list):
                errs.append(f"inputs_echo.{f}가 배열이 아님 — 형식오류(렌더 거부)")
        # core_scriptures 원소는 객체/문자열(ref), needs_care 원소는 문자열이어야 함(검증기 parity)
        for ci, cs in enumerate(_l(ie_g.get("core_scriptures"))):
            if isinstance(cs, list):
                errs.append(f"inputs_echo.core_scriptures[{ci}]가 배열 — ref 문자열/객체여야 함(렌더 거부)")
        for ni, nc in enumerate(_l(ie_g.get("needs_care"))):
            if isinstance(nc, (dict, list)):
                errs.append(f"inputs_echo.needs_care[{ni}]가 문자열이 아님(dict/list) — 형식오류(렌더 거부)")
    # keep_list/stop_list 원소는 문자열, warnings 원소는 객체여야 함(검증기 parity)
    for f in ("keep_list", "stop_list"):
        for ei, el in enumerate(_l(g.get(f))):
            if isinstance(el, (dict, list)):
                errs.append(f"{f}[{ei}]가 문자열이 아님(dict/list) — 형식오류(렌더 거부)")
    for ei, el in enumerate(_l(g.get("warnings"))):
        if not isinstance(el, dict):
            errs.append(f"warnings[{ei}]가 객체가 아님 — 형식오류(렌더 거부)")
    chosen_theme = _d(g.get("meta")).get("chosen_theme")
    if not chosen_theme or not str(chosen_theme).strip():
        errs.append("meta.chosen_theme 비어 있음 — 표어 먼저 확정(#2)")
    total = 0
    for ax in _l(g.get("axes")):
        if not isinstance(ax, dict):
            errs.append("axes 원소가 객체가 아님 — 형식오류(렌더 거부)")
            continue
        if "goals" in ax and not isinstance(ax.get("goals"), list):
            errs.append(f"'{ax.get('axis')}' 축 goals가 배열이 아님 — 형식오류(렌더 거부)")
        goals = _l(ax.get("goals"))
        if len(goals) > GOALS_PER_AXIS_MAX:
            errs.append(f"'{ax.get('axis')}' 축 목표 {len(goals)}개 — 축당 최대 {GOALS_PER_AXIS_MAX}(과욕)")
        total += len(goals)
        for j, goal in enumerate(goals):
            if not isinstance(goal, dict):
                errs.append(f"{ax.get('axis')} 목표[{j}] 형식오류 — 객체여야 함(렌더 거부)")
                continue
            dg = goal.get("direction_goal")
            if not (dg if isinstance(dg, str) else "").strip():
                errs.append(f"{ax.get('axis')} 목표[{j}] direction_goal 없음 — 환원(정성 누락)")
            inds = _l(goal.get("indicators"))
            if not inds:
                errs.append(f"{ax.get('axis')} 목표[{j}] indicators 없음 — 환원(정량 누락)")
            for k, ind in enumerate(inds):
                if not isinstance(ind, dict):
                    errs.append(f"{ax.get('axis')} 목표[{j}] 지표[{k}] 형식오류 — 객체여야 함(렌더 거부)")
                    continue
                if "baseline" in ind and not isinstance(ind.get("baseline"), dict):
                    errs.append(f"{ax.get('axis')} 목표[{j}] 지표[{k}].baseline 형식오류 — 객체여야 함(렌더 거부)")
    if total == 0:
        errs.append("핵심 목표 0개")
    if total > GOALS_TOTAL_MAX:
        errs.append(f"전체 목표 {total}개 — 최대 {GOALS_TOTAL_MAX}(과욕)")
    return errs


def flatten_rows(g):
    """축·목표·지표를 점검표 행으로 펼친다."""
    rows = []
    for ax in _ld(g.get("axes")):
        name = str(ax.get("axis", ""))
        for goal in _ld(ax.get("goals")):
            dg = str(goal.get("direction_goal", "") or "")
            rat = str(goal.get("rationale", "") or "")
            inds = _ld(goal.get("indicators"))
            if not inds:
                rows.append([name, dg, "", "", "", "", rat])
                continue
            for ind in inds:
                base = _d(ind.get("baseline"))
                baseline_txt = base.get("value")
                rng = base.get("range")
                if baseline_txt is None and isinstance(rng, list) and len(rng) >= 2:
                    baseline_txt = f"{rng[0]}~{rng[1]}"
                conf = base.get("confidence", "")
                baseline_disp = f"{baseline_txt}" + (f" (신뢰{conf})" if conf else "")
                signal = " ·신호" if ind.get("signal_only") else ""
                rows.append([
                    name, dg, str(ind.get("label", "")) + signal,
                    baseline_disp, str(ind.get("target", "")), str(ind.get("by", "")),
                    rat,
                ])
    return rows


def build_md(g, rows):
    meta = _d(g.get("meta"))
    lines = []
    lines.append(f"# 2027 3축 목표·지표 점검표")
    lines.append("")
    lines.append(f"- 표어: **{meta.get('chosen_theme', '(미확정)')}**")
    lines.append(f"- 상태: {meta.get('status', '제안')}"
                 + ("  ·  ✅ 목사 확정" if g.get("confirmed") else "  ·  🙏 목사 확정 대기"))
    decl = _d(g.get("inputs_echo")).get("direction_declaration")
    if decl:
        lines.append(f"- 방향: {decl}")
    lines.append("")
    lines.append("## 3축 목표·지표")
    lines.append("")
    lines.append("| 축 | 방향 목표 | 신호 지표 | 기준선 | 올해 목표치 | 기한 |")
    lines.append("|---|---|---|---|---|---|")
    for r in rows:
        lines.append(f"| {r[0]} | {r[1]} | {r[2]} | {r[3]} | {r[4]} | {r[5]} |")
    lines.append("")

    # 유지/중단은 문자열 목록 — dict/list 원소는 산출물 오염 방지 위해 제외
    keep = [str(k) for k in _l(g.get("keep_list")) if not isinstance(k, (dict, list))]
    stop = [str(s) for s in _l(g.get("stop_list")) if not isinstance(s, (dict, list))]
    lines.append("## 올해 유지 / 올해는 하지 않을 것")
    lines.append("")
    lines.append("**유지 (지킬 것):**")
    for k in keep:
        lines.append(f"- {k}")
    if not keep:
        lines.append("- (없음)")
    lines.append("")
    lines.append("**중단 (올해는 하지 않을 것 — 더하기 전 빼기):**")
    for s in stop:
        lines.append(f"- {s}")
    if not stop:
        lines.append("- (없음)")
    lines.append("")

    warns = _ld(g.get("warnings"))
    if warns:
        lines.append("## ⚖️ 짐 가늠 (과욕·자원 경고)")
        lines.append("")
        for w in warns:
            lines.append(f"- **[{w.get('type', '')}]** {w.get('text', '')}  _(근거: {w.get('basis', '')})_")
        lines.append("")

    lines.append("> 지표는 성과의 채찍이 아니라 돌봄이 닿지 않은 양을 찾는 목자의 눈입니다.")
    lines.append("> 분기 점검(#9 review-rhythm)에서 이 표의 목표치를 신호등으로 따라갑니다.")
    return "\n".join(lines) + "\n"


XLSX_HEADER = ["축", "방향 목표", "신호 지표", "기준선", "올해 목표치", "기한",
               "Q1", "Q2", "Q3", "Q4", "근거/비고"]


def build_xlsx(g, rows, out_path):
    """openpyxl 있으면 xlsx, 없으면 csv 폴백. 실제 쓴 경로 반환."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        csv_path = out_path.rsplit(".", 1)[0] + ".csv"
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(XLSX_HEADER)
            for r in rows:
                w.writerow([r[0], r[1], r[2], r[3], r[4], r[5], "", "", "", "", r[6]])
            w.writerow([])
            w.writerow(["[올해 유지]"])
            for k in _l(g.get("keep_list")):
                if isinstance(k, (dict, list)):
                    continue
                w.writerow([str(k)])
            w.writerow(["[올해는 하지 않을 것]"])
            for s in _l(g.get("stop_list")):
                if isinstance(s, (dict, list)):
                    continue
                w.writerow([str(s)])
        return csv_path

    wb = Workbook()
    ws = wb.active
    ws.title = "3축 목표 점검표"
    head_fill = PatternFill("solid", fgColor="594B3B")
    head_font = Font(bold=True, color="FFFFFF")
    ws.append(XLSX_HEADER)
    for c in ws[1]:
        c.fill = head_fill; c.font = head_font; c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([r[0], r[1], r[2], r[3], r[4], r[5], "", "", "", "", r[6]])
    widths = [8, 30, 22, 16, 16, 8, 6, 6, 6, 6, 30]
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = wd
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    ws2 = wb.create_sheet("유지·중단")
    ws2.append(["구분", "항목"])
    for c in ws2[1]:
        c.fill = head_fill; c.font = head_font
    for k in _l(g.get("keep_list")):
        if isinstance(k, (dict, list)):
            continue
        ws2.append(["유지", str(k)])
    for s in _l(g.get("stop_list")):
        if isinstance(s, (dict, list)):
            continue
        ws2.append(["중단", str(s)])
    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 50

    warns = _ld(g.get("warnings"))
    if warns:
        ws3 = wb.create_sheet("짐 가늠 경고")
        ws3.append(["유형", "내용", "근거"])
        for c in ws3[1]:
            c.fill = head_fill; c.font = head_font
        for w in warns:
            ws3.append([str(w.get("type", "")), str(w.get("text", "")), str(w.get("basis", ""))])
        for col, wd in zip("ABC", [10, 40, 30]):
            ws3.column_dimensions[col].width = wd

    wb.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 2:
        print("사용: python3 build_checklist.py <goals.json> [출력디렉토리]"); sys.exit(2)
    src = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(src))
    try:
        with open(src, encoding="utf-8") as f:
            g = json.load(f)
    except FileNotFoundError:
        print(f"[오류] 파일 없음: {src}"); sys.exit(2)
    except json.JSONDecodeError as e:
        print(f"[오류] JSON 파싱 실패: {e}"); sys.exit(2)

    if not isinstance(g, dict):
        print("[오류] 최상위가 객체(dict)가 아님 — goals.json은 객체여야 함"); sys.exit(2)

    errs = gate(g)
    if errs:
        print("=" * 60)
        print("렌더 거부 — 게이트 위반 (fidelity-guard 먼저 통과시키라)")
        print("=" * 60)
        for e in errs:
            print(f"  ✗ {e}")
        sys.exit(1)

    os.makedirs(out_dir, exist_ok=True)
    rows = flatten_rows(g)

    md_path = os.path.join(out_dir, "goal-checklist.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(build_md(g, rows))

    xlsx_path = build_xlsx(g, rows, os.path.join(out_dir, "goal-checklist.xlsx"))

    print("렌더 완료 ✓")
    print(f"  · {md_path}")
    print(f"  · {xlsx_path}" + ("  (openpyxl 없어 csv 폴백)" if xlsx_path.endswith(".csv") else ""))
    sys.exit(0)


if __name__ == "__main__":
    main()
