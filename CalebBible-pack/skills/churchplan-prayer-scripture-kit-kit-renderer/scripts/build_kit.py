#!/usr/bin/env python3
"""
build_kit.py — churchplan-prayer-scripture-kit 영성 키트 렌더러(자기완결).

prayer-kit.json → ① prayer-kit.docx (통독표 + 월별 기도제목 + 가정예배 가이드 + 새벽 안내문)
                  ② prayer-kit.pdf (weasyprint)

렌더 전 핵심 게이트를 재강제한다(cys-insight 패턴):
  필수 블록 · status · 통독 은혜 융통성(grace_frame) · 율법 어휘 없음
  · 성경 실재(bible_validator — 통독 portions·가정예배/새벽 scripture의 장·절·장범위 끝까지).
  위반 시 렌더 거부(종료코드 1) — 가짜·범위밖 본문이 docx/pdf로 먼저 인쇄되는 것을 차단.
곡해·아전인수 등 내용 정밀 검증은 fidelity-guard의 validate_kit.py가 사람 눈과 함께 맡는다.

사용:  python3 build_kit.py <prayer-kit.json> [출력디렉토리]
종료코드: 0 = 성공, 1 = 게이트 위반(렌더 거부), 2 = 파일/JSON 오류
라이브러리 없으면 폴백(python-docx→md·weasyprint→pdf 생략).
"""
import html
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
try:
    import bible_validator as bv  # noqa: E402  같은 폴더 정경 데이터(이식·diff 0)
except ImportError:
    bv = None

FONT = "나눔명조"
SEPIA = "594B3B"
# 가드레일/스키마 단일 진실이 열거한 금지 어휘. '반드시'는 단독 토큰(검증기 parity — '반드시 매일'
# bigram은 '반드시 통독을 마치십시오' 단독 강제를 놓쳤다). '매일 …해야'는 DAILY_OBLIG 근접 규칙이 맡는다.
LEGALISM = ["완주 못하면", "해야만", "해야 한다", "의무", "죄책", "게으름", "빠지면 안", "낙오", "반드시"]
DAILY_OBLIG = r"매일[가-힣\s]{0,14}?[가-힣]야(?:만|\s*(?:하|합|한|해|된|돼))"
# 부정·면제 표지 — 어휘가 *뒤집힌* 맥락일 때만 면제. '아닙니다'는 '아니'로 분절되지 않으므로 '아닙'도 포함.
# 무관한 위치의 '없/않/뿐' 단독으로는 면제하지 않는다(문자열 전체 면제 차단).
NEGATE = r"(?:아니|아닙|않|말고)"
NEGATE_WINDOW = 12
# 다어절 토큰의 리터럴 공백을 \s*로 — 이중공백·무공백·탭 변형 우회 차단(validate_kit.py와 동일·parity).
_LEGALISM_RE = [(w, re.compile(re.escape(w).replace("\\ ", r"\s*"))) for w in LEGALISM]


def _legalism_hits(text):
    """어휘 직후 윈도 안에 부정 표지가 있으면(뒤집힌 맥락) 그 어휘만 면제."""
    if not isinstance(text, str) or not text.strip():
        return []
    hits = []
    for w, pat in _LEGALISM_RE:
        for m in pat.finditer(text):
            window = text[m.end():m.end() + NEGATE_WINDOW]
            if re.search(NEGATE, window):
                continue
            hits.append(w)
            break
    # '매일 …해야' 근접 패턴(검증기와 동일) — 의무 활용 직후 부정 표지 없으면 위반.
    for m in re.finditer(DAILY_OBLIG, text):
        window = text[m.end():m.end() + NEGATE_WINDOW]
        if re.search(NEGATE, window):
            continue
        hits.append("매일 …해야")
        break
    return hits


def check_scripture(text):
    """본문 표기가 정경 실재인지 — 책·장·절·장범위 끝까지(validate_kit.check_scripture와 동일 로직).
    (True, info) | (False, 사유). bv 없으면 (True, None)로 통과(폴백)."""
    if bv is None:
        return (True, None)
    raw = str(text).strip()
    if not raw:
        return (False, raw)
    m = re.match(r"^\s*([가-힣]+)\s*(.*)$", raw)
    if not m:
        return (False, raw)
    book = bv.normalize_book(m.group(1).strip())
    if not book:
        return (False, raw)
    rest = m.group(2).strip()
    if not rest:
        return (True, book)
    if ":" in rest:
        ok, msg = bv.validate_reference(raw)
        return (ok, raw if ok else msg)
    # 장 표기는 '장'·'편'(시편) 모두 허용 — 시편 본문은 관례상 '편'을 쓴다(카탈로그 §1 "시편 150편").
    nch = len(bv.BOOK_CHAPTERS[book])
    rng = re.match(r"^(\d+)\s*(?:장|편)?\s*(?:[-–~]\s*(\d+)\s*(?:장|편)?)?$", rest)
    if not rng:
        return (False, raw)
    start = int(rng.group(1))
    end = int(rng.group(2)) if rng.group(2) else start
    if start < 1 or start > nch or end < start or end > nch:
        return (False, f"{raw} — {book}는 총 {nch}장. 범위 밖")
    return (True, book)


def gate(p):
    """렌더 전 핵심 게이트(성경 실재 포함 — 가짜 본문이 인쇄물에 먼저 들어가는 것 차단)."""
    errs = []
    if not isinstance(p, dict):
        return [f"최상위 JSON이 객체(dict)가 아님: {type(p).__name__}"]
    for blk in ("meta", "reading_track", "prayer_topics", "family_worship", "dawn_prayer_season"):
        if blk not in p:
            errs.append(f"필수 블록 '{blk}' 없음")
        elif not isinstance(p.get(blk), dict):
            errs.append(f"필수 블록 '{blk}'가 dict가 아님: {type(p.get(blk)).__name__} (truthy 비-dict)")
    meta = _block(p, "meta")
    if meta.get("status") not in ("제안", "확정"):
        errs.append("meta.status 가 제안|확정 이 아님")
    rt = _block(p, "reading_track")
    if not rt.get("grace_frame"):
        errs.append("reading_track.grace_frame 없음 — 통독 은혜 융통성(율법화 차단)")
    # 율법 어휘
    for s in _collect(p):
        for w in _legalism_hits(s):
            errs.append(f"율법 어휘 '{w}' — 은혜의 초청: \"{s[:24]}\"")
    # 성경 실재 — 가짜·범위밖 본문이 docx/pdf로 먼저 인쇄되지 않게 렌더 단계에서도 차단
    for d in _dicts(rt.get("schedule")):
        for portion in (d.get("portions") if isinstance(d.get("portions"), list) else []):
            ok, info = check_scripture(portion)
            if not ok:
                errs.append(f"가짜성경 reading_track.schedule '{portion}' — {info}")
    for s in _dicts(_block(p, "family_worship").get("samples")):
        if s.get("scripture"):
            ok, info = check_scripture(s.get("scripture"))
            if not ok:
                errs.append(f"가짜성경 family_worship.samples '{s.get('scripture')}' — {info}")
    for s in _dicts(_block(p, "dawn_prayer_season").get("seasons")):
        if s.get("scripture"):
            ok, info = check_scripture(s.get("scripture"))
            if not ok:
                errs.append(f"가짜성경 dawn_prayer_season.seasons '{s.get('scripture')}' — {info}")
    return errs


def _collect(obj):
    out = []
    if isinstance(obj, dict):
        for v in obj.values():
            out += _collect(v)
    elif isinstance(obj, list):
        for v in obj:
            out += _collect(v)
    elif isinstance(obj, str):
        out.append(obj)
    return out


def _block(p, name):
    """최상위 블록을 dict로 안전 추출(문자열·리스트 등 형오류면 빈 dict)."""
    v = p.get(name, {})
    return v if isinstance(v, dict) else {}


def _dicts(value):
    """value가 list면 그 안의 dict 원소만, 아니면 빈 리스트(적대적 형 방어)."""
    if not isinstance(value, list):
        return []
    return [it for it in value if isinstance(it, dict)]


def _scalars(value):
    """value가 list면 스칼라 원소(dict·list 제외)만, 아니면 빈 리스트.
    문자열 리스트 필드(order·topics·reflection_questions)에 섞인 dict/list가 산출물에 새는 것을 차단."""
    if not isinstance(value, list):
        return []
    return [it for it in value if not isinstance(it, (dict, list))]


def _strjoin(value, sep):
    """value가 list면 스칼라 원소만 문자열로 안전 결합, 아니면 빈 문자열.
    dict·list 원소는 산출물 오염 차단을 위해 제외(문자열 리스트 필드에 섞인 비-스칼라 방어)."""
    if not isinstance(value, list):
        return ""
    return sep.join(str(v) for v in value if not isinstance(v, (dict, list)))


def esc(x):
    return html.escape(str(x if x is not None else ""))


def kit_html(p):
    meta = _block(p, "meta")
    rt = _block(p, "reading_track")
    pt = _block(p, "prayer_topics")
    fw = _block(p, "family_worship")
    dp = _block(p, "dawn_prayer_season")
    theme = meta.get("chosen_theme", "")
    L = ["<!DOCTYPE html><html lang='ko'><head><meta charset='utf-8'>",
         "<title>2027 영성 실천 키트</title>",
         f"""<style>
@page{{size:A4;margin:20mm 18mm;@bottom-center{{content:counter(page);color:#9a9a9a;font-size:9pt}}}}
body{{font-family:'{FONT}',serif;color:#241f1a;line-height:1.7;font-size:11pt}}
h1{{color:#{SEPIA};font-size:21pt;text-align:center;margin:0 0 3pt}}
.theme{{text-align:center;color:#{SEPIA};font-size:12.5pt;margin:0 0 4pt}}
.grace{{background:#F3EFE6;border-radius:7pt;padding:9pt 13pt;color:#5e5446;font-size:10.5pt;margin:10pt 0 16pt;text-align:center}}
h2{{color:#{SEPIA};font-size:14pt;border-bottom:1.5pt solid #{SEPIA};padding-bottom:3pt;margin:18pt 0 7pt}}
table{{border-collapse:collapse;width:100%;font-size:10.5pt;margin:5pt 0}}
th,td{{border:1px solid #E0D7C8;padding:5pt 7pt;text-align:left;vertical-align:top}}
th{{background:#{SEPIA};color:#fff}}
.chk{{width:18pt;text-align:center;color:#b59a6a}}
.q{{color:#5e5446;font-size:10.5pt}}
.src{{color:#a99;font-size:8.5pt}}
ul{{margin:3pt 0;padding-left:18pt}}
</style></head><body>"""]
    L.append("<h1>2027 영성 실천 키트</h1>")
    if theme:
        L.append(f"<p class='theme'>{esc(theme)}</p>")

    # 통독 트랙
    L.append("<h2>📖 연간 성경통독 트랙</h2>")
    L.append(f"<p>트랙: <b>{esc(rt.get('track',''))}</b> &nbsp; <span class='src'>{esc(rt.get('track_basis',''))}</span></p>")
    if rt.get("grace_frame"):
        L.append(f"<div class='grace'>🕊️ {esc(rt.get('grace_frame'))}</div>")
    sched = _dicts(rt.get("schedule"))
    if sched:
        L.append("<table><tr><th>구간</th><th>본문</th>" + ("<th class='chk'>✓</th>" if rt.get("checkbox") else "") + "</tr>")
        for d in sched:
            chk = "<td class='chk'>☐</td>" if rt.get("checkbox") else ""
            L.append(f"<tr><td>{esc(d.get('unit',''))}</td><td>{esc(_strjoin(d.get('portions'), ' · '))}</td>{chk}</tr>")
        L.append("</table>")
    if rt.get("serves_indicator"):
        L.append(f"<p class='src'>※ {esc(rt.get('serves_indicator'))}</p>")

    # 월별 기도제목
    L.append("<h2>🙏 월별 전교인 기도제목</h2>")
    if pt.get("frame"):
        L.append(f"<p class='q'>{esc(pt.get('frame'))}</p>")
    L.append("<table><tr><th>월</th><th>기도제목</th></tr>")
    for m_ in _dicts(pt.get("months")):
        tlist = _scalars(m_.get("topics"))
        topics = "<br>".join(f"· {esc(t)}" for t in tlist)
        L.append(f"<tr><td>{esc(m_.get('month',''))}월</td><td>{topics}</td></tr>")
    L.append("</table>")

    # 가정예배
    L.append("<h2>🏠 가정예배 가이드 (함께 나누는 1:1 말씀나눔)</h2>")
    if fw.get("frame"):
        L.append(f"<p class='q'>{esc(fw.get('frame'))}</p>")
    order = _scalars(fw.get("order"))
    if order:
        L.append("<p><b>순서</b>: " + " → ".join(esc(o) for o in order) + "</p>")
    for s in _dicts(fw.get("samples")):
        L.append(f"<p><b>{esc(s.get('scripture',''))}</b></p>")
        qs = _scalars(s.get("reflection_questions"))
        if qs:
            L.append("<ul>" + "".join(f"<li class='q'>{esc(q)}</li>" for q in qs) + "</ul>")
        if s.get("prayer"):
            L.append(f"<p class='q'>기도: {esc(s.get('prayer'))}</p>")

    # 새벽 시즌
    L.append("<h2>🌅 새벽기도 시즌 안내</h2>")
    if dp.get("frame"):
        L.append(f"<p class='q'>{esc(dp.get('frame'))}</p>")
    for s in _dicts(dp.get("seasons")):
        L.append(f"<p><b>{esc(s.get('name',''))}</b> ({esc(s.get('when',''))}) — {esc(s.get('theme',''))} "
                 f"<span class='src'>{esc(s.get('scripture',''))}</span></p>")
    if dp.get("serves_indicator"):
        L.append(f"<p class='src'>※ {esc(dp.get('serves_indicator'))}</p>")

    L.append("<p class='src' style='margin-top:16pt;text-align:center'>은혜의 초청입니다. 못 따라가도 괜찮습니다 — 언제든 다시 함께 걸으면 됩니다.</p>")
    L.append("</body></html>")
    return "\n".join(L)


def build_docx(p, path):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
    except ImportError:
        return None
    meta = _block(p, "meta")
    rt = _block(p, "reading_track")
    pt = _block(p, "prayer_topics")
    fw = _block(p, "family_worship")
    dp = _block(p, "dawn_prayer_season")
    sepia = RGBColor(0x59, 0x4B, 0x3B)
    muted = RGBColor(0x8A, 0x7F, 0x70)

    def setfont(run, size, color=None, bold=False):
        run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
        if color:
            run.font.color.rgb = color
        rPr = run._element.get_or_add_rPr()
        rF = rPr.find(qn("w:rFonts"))
        if rF is None:
            rF = rPr.makeelement(qn("w:rFonts"), {}); rPr.append(rF)
        for a in ("w:eastAsia", "w:ascii", "w:hAnsi"):
            rF.set(qn(a), FONT)

    def heading(txt):
        h = doc.add_paragraph(); setfont(h.add_run(txt), 14, sepia, True)

    doc = Document()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    setfont(t.add_run("2027 영성 실천 키트"), 20, sepia, True)
    if meta.get("chosen_theme"):
        th = doc.add_paragraph(); th.alignment = WD_ALIGN_PARAGRAPH.CENTER
        setfont(th.add_run(meta.get("chosen_theme")), 12.5, sepia)

    heading("📖 연간 성경통독 트랙")
    tr = doc.add_paragraph(); setfont(tr.add_run(f"트랙: {rt.get('track','')}  ({rt.get('track_basis','')})"), 11)
    if rt.get("grace_frame"):
        g = doc.add_paragraph(); setfont(g.add_run(f"🕊️ {rt.get('grace_frame')}"), 10.5, muted)
    for d in _dicts(rt.get("schedule")):
        line = f"{d.get('unit','')}: {_strjoin(d.get('portions'), ' · ')}" + ("  ☐" if rt.get("checkbox") else "")
        bp = doc.add_paragraph(); setfont(bp.add_run(line), 10.5)
    if rt.get("serves_indicator"):
        s = doc.add_paragraph(); setfont(s.add_run(f"※ {rt.get('serves_indicator')}"), 8.5, muted)

    heading("🙏 월별 전교인 기도제목")
    if pt.get("frame"):
        f = doc.add_paragraph(); setfont(f.add_run(pt.get("frame")), 10.5, muted)
    for m_ in _dicts(pt.get("months")):
        line = f"{m_.get('month','')}월: " + _strjoin(m_.get("topics"), " / ")
        bp = doc.add_paragraph(); setfont(bp.add_run(line), 10.5)

    heading("🏠 가정예배 가이드 (함께 나누는 1:1 말씀나눔)")
    if fw.get("frame"):
        f = doc.add_paragraph(); setfont(f.add_run(fw.get("frame")), 10.5, muted)
    if isinstance(fw.get("order"), list) and fw.get("order"):
        o = doc.add_paragraph(); setfont(o.add_run("순서: " + _strjoin(fw.get("order"), " → ")), 10.5)
    for s in _dicts(fw.get("samples")):
        sp = doc.add_paragraph(); setfont(sp.add_run(s.get("scripture", "")), 11, bold=True)
        for q in _scalars(s.get("reflection_questions")):
            qp = doc.add_paragraph(); setfont(qp.add_run(f"  · {q}"), 10.5, muted)
        if s.get("prayer"):
            pp = doc.add_paragraph(); setfont(pp.add_run(f"  기도: {s.get('prayer')}"), 10.5, muted)

    heading("🌅 새벽기도 시즌 안내")
    if dp.get("frame"):
        f = doc.add_paragraph(); setfont(f.add_run(dp.get("frame")), 10.5, muted)
    for s in _dicts(dp.get("seasons")):
        line = f"{s.get('name','')} ({s.get('when','')}) — {s.get('theme','')}  {s.get('scripture','')}"
        bp = doc.add_paragraph(); setfont(bp.add_run(line), 10.5)

    end = doc.add_paragraph(); end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    setfont(end.add_run("은혜의 초청입니다. 못 따라가도 괜찮습니다 — 언제든 다시 함께 걸으면 됩니다."), 9, muted)
    doc.save(path)
    return path


def kit_md(p):
    meta = _block(p, "meta")
    rt = _block(p, "reading_track")
    pt = _block(p, "prayer_topics")
    fw = _block(p, "family_worship")
    dp = _block(p, "dawn_prayer_season")
    L = ["# 2027 영성 실천 키트", ""]
    if meta.get("chosen_theme"):
        L += [f"**{meta.get('chosen_theme')}**", ""]
    L += ["## 📖 연간 성경통독 트랙", "", f"트랙: {rt.get('track','')} ({rt.get('track_basis','')})", ""]
    if rt.get("grace_frame"):
        L += [f"> 🕊️ {rt.get('grace_frame')}", ""]
    for d in _dicts(rt.get("schedule")):
        L.append(f"- {d.get('unit','')}: {_strjoin(d.get('portions'), ' · ')}")
    L += ["", "## 🙏 월별 전교인 기도제목", ""]
    for m_ in _dicts(pt.get("months")):
        L.append(f"- {m_.get('month','')}월: " + _strjoin(m_.get("topics"), " / "))
    L += ["", "## 🏠 가정예배 가이드", ""]
    if isinstance(fw.get("order"), list) and fw.get("order"):
        L.append("순서: " + _strjoin(fw.get("order"), " → "))
    for s in _dicts(fw.get("samples")):
        rq = _scalars(s.get("reflection_questions"))
        L += [f"- **{s.get('scripture','')}**"] + [f"  - {q}" for q in rq]
    L += ["", "## 🌅 새벽기도 시즌", ""]
    for s in _dicts(dp.get("seasons")):
        L.append(f"- {s.get('name','')} ({s.get('when','')}) — {s.get('theme','')} {s.get('scripture','')}")
    return "\n".join(L) + "\n"


def render_pdf(html_str, path):
    try:
        from weasyprint import HTML
        HTML(string=html_str).write_pdf(path)
        return path
    except Exception:
        return None


def main():
    if len(sys.argv) < 2:
        print("사용: python3 build_kit.py <prayer-kit.json> [출력디렉토리]"); sys.exit(2)
    src = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(src))
    try:
        with open(src, encoding="utf-8") as f:
            p = json.load(f)
    except FileNotFoundError:
        print(f"[오류] 파일 없음: {src}"); sys.exit(2)
    except json.JSONDecodeError as ex:
        print(f"[오류] JSON 파싱 실패: {ex}"); sys.exit(2)

    errs = gate(p)
    if errs:
        print("=" * 60)
        print("렌더 거부 — 게이트 위반 (해당 분석으로 되돌려 보완하라)")
        print("=" * 60)
        for er in errs:
            print(f"  ✗ {er}")
        sys.exit(1)

    os.makedirs(out_dir, exist_ok=True)
    made = []
    docx_path = build_docx(p, os.path.join(out_dir, "prayer-kit.docx"))
    if docx_path:
        made.append(docx_path)
    else:
        mp = os.path.join(out_dir, "prayer-kit.md")
        with open(mp, "w", encoding="utf-8") as f:
            f.write(kit_md(p))
        made.append(mp + "  (python-docx 없어 md 폴백)")
    pdf = render_pdf(kit_html(p), os.path.join(out_dir, "prayer-kit.pdf"))
    made.append(pdf if pdf else "prayer-kit.pdf 생략(weasyprint 없음)")

    print("렌더 완료 ✓")
    for m in made:
        print(f"  · {m}")
    sys.exit(0)


if __name__ == "__main__":
    main()
