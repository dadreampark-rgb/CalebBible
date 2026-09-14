#!/usr/bin/env python3
"""
build_action.py — churchplan-action-planner 부서별 실행계획표 렌더러.

action.json → ① action-plan.md (부서별 실행계획표 + 공백 목록 + 과잉·중복 경고)
              ② action-plan.xlsx (부서별 실행 시트 + 공백·경고 시트, openpyxl 없으면 .csv)

렌더 전 핵심 게이트를 재강제한다(cys-insight 패턴):
  필수 블록 · 표어 · 부서 · 각 action source_goal · status. 위반 시 렌더 거부(종료코드 1).
전체 정밀 검증(근거 추적·공백 정직·캘린더 정합)은 fidelity-guard의 validate_action.py가 맡는다.

사용:  python3 build_action.py <action.json> [출력디렉토리]
종료코드: 0 = 성공, 1 = 게이트 위반(렌더 거부), 2 = 파일/JSON 오류
표준 라이브러리만 사용(xlsx는 openpyxl 있을 때만, 없으면 csv).
"""
import csv
import json
import os
import sys


def _d(x):
    return x if isinstance(x, dict) else {}


def _l(x):
    return x if isinstance(x, list) else []


def _ld(x):
    return [v for v in x if isinstance(v, dict)] if isinstance(x, list) else []


def gate(a):
    """렌더 전 핵심 게이트."""
    errs = []
    for blk in ("meta", "actions"):
        if blk not in a:
            errs.append(f"필수 블록 '{blk}' 없음")
    # 스키마 필드 형식 오염 검사(검증기-렌더러 parity, 조용한 통과 금지)
    for k in ("meta", "inputs_echo", "department_table"):
        if k in a and not isinstance(a.get(k), dict):
            errs.append(f"'{k}' 가 객체(dict)가 아님 — 스키마 필드 오염")
    for k in ("actions", "gaps", "warnings"):
        if k in a and not isinstance(a.get(k), list):
            errs.append(f"'{k}' 가 리스트가 아님 — 스키마 필드 오염")
    _ie = a.get("inputs_echo")
    if isinstance(_ie, dict):
        if "from_goals" in _ie and not isinstance(_ie.get("from_goals"), list):
            errs.append("inputs_echo.from_goals 가 리스트가 아님 — 목표 출처 오염")
        if "from_calendar" in _ie and not isinstance(_ie.get("from_calendar"), dict):
            errs.append("inputs_echo.from_calendar 가 객체(dict)가 아님 — 캘린더 출처 오염")

    # list-of-dict 필드 원소 오염 검사(검증기-렌더러 parity)
    def _bad_dict_elems(lst, label):
        if isinstance(lst, list):
            for j, el in enumerate(lst):
                if not isinstance(el, dict):
                    errs.append(f"{label}[{j}] 객체(dict) 아님 — 리스트 원소 오염")
    _bad_dict_elems(_d(_ie).get("from_goals"), "inputs_echo.from_goals")
    _bad_dict_elems(_d(_d(_ie).get("from_calendar")).get("event_weeks"), "inputs_echo.from_calendar.event_weeks")
    _bad_dict_elems(a.get("gaps"), "gaps")
    _bad_dict_elems(a.get("warnings"), "warnings")

    meta = _d(a.get("meta"))
    if not str(meta.get("chosen_theme") or "").strip():
        errs.append("meta.chosen_theme 비어 있음 — 표어 먼저 확정")
    if not _l(meta.get("departments")):
        errs.append("meta.departments 비어 있음 — 부서 구조 먼저 입력")
    if meta.get("status") not in ("제안", "확정"):
        errs.append("meta.status 가 제안|확정 이 아님")
    acts = _l(a.get("actions"))
    if len(acts) == 0:
        errs.append("actions 0개 — 분해할 행동 없음")
    for i, act in enumerate(acts):
        if not isinstance(act, dict):
            errs.append(f"actions[{i}] 객체 아님 — 렌더 불가"); continue
        if not act.get("source_goal"):
            errs.append(f"actions[{act.get('id', i)}] source_goal 없음 — 근거 없는 행동")
        _p = act.get("purpose")
        if not _p or not str(_p).strip():
            errs.append(f"actions[{act.get('id', i)}] purpose 없음 — 목적은 필수(schema 필수 블록)")
        _aid = act.get("id", i)
        if "milestone" in act and not isinstance(act.get("milestone"), dict):
            errs.append(f"actions[{_aid}].milestone 객체(dict) 아님 — 필드 오염")
        for k in ("resources", "flags", "monthly_steps"):
            if k in act and not isinstance(act.get(k), list):
                errs.append(f"actions[{_aid}].{k} 리스트 아님 — 필드 오염")
        _bad_dict_elems(act.get("monthly_steps"), f"actions[{_aid}].monthly_steps")
        for k in ("resources", "flags"):
            if isinstance(act.get(k), list):
                for j, el in enumerate(act.get(k)):
                    if isinstance(el, (dict, list)):
                        errs.append(f"actions[{_aid}].{k}[{j}] dict/list — 스칼라만 허용(원소 오염)")
    return errs


def owner_disp(act):
    o = act.get("owner")
    if (not o) or str(o).strip().upper() == "TBD":
        return "⚠ TBD(미정)"
    return str(o)


def res_disp(act):
    # 비-리스트는 미정 처리, 리스트 안 dict/list 원소는 산출물 오염 차단(제외)
    r = [v for v in _l(act.get("resources")) if not isinstance(v, (dict, list))]
    return ", ".join(str(v) for v in r) if r else "⚠ 미정"


def build_md(a):
    meta = _d(a.get("meta"))
    raw_actions = _ld(a.get("actions"))
    actions = {act.get("id"): act for act in raw_actions}
    dept_table = _d(a.get("department_table"))
    depts = _l(meta.get("departments")) or list(dept_table.keys())

    L = []
    L.append("# 2027 부서별 실행계획표")
    L.append("")
    L.append(f"- 표어: **{meta.get('chosen_theme', '(미확정)')}**")
    L.append(f"- 상태: " + ("✅ 목사 확정" if a.get("confirmed") else "🙏 목사 확정 대기"))
    L.append("")

    # 부서별 표
    rendered_ids = set()
    for d in depts:
        ids = _l(dept_table.get(d))
        if not ids:
            # department_table에 없으면 actions에서 직접 수집
            ids = [act.get("id") for act in raw_actions if act.get("department") == d]
        if not ids:
            continue
        L.append(f"## [{str(d)}]")
        L.append("")
        L.append("| 행동 | 목적 | 시기 | 담당 | 자원 | 마일스톤 |")
        L.append("|---|---|---|---|---|---|")
        for aid in ids:
            act = actions.get(aid)
            if not act:
                continue
            rendered_ids.add(aid)
            ms = _d(act.get("milestone"))
            L.append(f"| {act.get('source_goal','')} | {act.get('purpose','')} | {act.get('timing','')} "
                     f"| {owner_disp(act)} | {res_disp(act)} | {ms.get('label','')} ({ms.get('by','')}) |")
        L.append("")
        # 월별 단계(있는 것만)
        for aid in ids:
            act = actions.get(aid)
            steps = _ld((act or {}).get("monthly_steps"))
            if steps:
                seq = " → ".join(f"{s.get('month','')}월 {s.get('step','')}" for s in steps)
                L.append(f"- {(act or {}).get('purpose','')}: {seq}")
        L.append("")

    # 고아 action (departments 어디에도 배치되지 않음 — 부서 위장/오타로 표에서 누락)
    orphans = [act for act in raw_actions if act.get("id") not in rendered_ids]
    if orphans:
        L.append("## ⚠️ 부서 미배치 action (departments에 없는 부서 — fidelity-guard로 되돌릴 것)")
        L.append("")
        for act in orphans:
            L.append(f"- **{act.get('id','')}** [{act.get('department','(부서없음)')}] "
                     f"{act.get('source_goal','')} — meta.departments에 없는 부서라 부서별 표에서 누락됨")
        L.append("")

    # 공백 목록
    gaps = _ld(a.get("gaps"))
    L.append("## 🕳️ \"언제·누가·무엇으로\" 공백 (실행 전 채울 것)")
    L.append("")
    if gaps:
        for g in gaps:
            L.append(f"- **[{g.get('type','')}]** {g.get('text','')} _(행동 {g.get('action_id','')})_")
    else:
        L.append("- (공백 없음 — 모든 행동에 담당·시기·자원이 채워짐)")
    L.append("")

    # 경고
    warns = _ld(a.get("warnings"))
    if warns:
        L.append("## ⚖️ 과잉·중복 경고")
        L.append("")
        for w in warns:
            L.append(f"- **[{w.get('type','')}]** {w.get('text','')}  _(근거: {w.get('basis','')})_")
        L.append("")

    L.append("> 계획이 종이로 남지 않게, '언제·누가·무엇으로'의 공백을 정직하게 드러냅니다. 빈칸은 가짜로 메우지 않고, 한 손에 일이 몰리지 않게 비춥니다.")
    return "\n".join(L) + "\n"


GRID_HEADER = ["부서", "행동(목표)", "목적", "시기", "담당", "자원", "마일스톤", "기한", "공백"]


def grid_rows(a):
    rows = []
    for act in _ld(a.get("actions")):
        ms = _d(act.get("milestone"))
        # dict/list 원소는 join 오염 차단(제외), 나머지는 str 강제
        flags = [str(f) for f in _l(act.get("flags")) if not isinstance(f, (dict, list))]
        rows.append([
            act.get("department", ""), act.get("source_goal", ""), act.get("purpose", ""),
            act.get("timing", ""), owner_disp(act), res_disp(act),
            ms.get("label", ""), ms.get("by", ""), ", ".join(flags),
        ])
    rows.sort(key=lambda r: str(r[0]))  # 부서별 정렬(비-문자열 부서 방어)
    return rows


def build_xlsx(a, out_path):
    rows = grid_rows(a)
    gaps = _ld(a.get("gaps"))
    warns = _ld(a.get("warnings"))
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        csv_path = out_path.rsplit(".", 1)[0] + ".csv"
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(GRID_HEADER)
            for r in rows:
                w.writerow(r)
            w.writerow([])
            w.writerow(["[공백 목록]"])
            for g in gaps:
                w.writerow([g.get("action_id", ""), g.get("type", ""), g.get("text", "")])
            w.writerow([])
            w.writerow(["[경고]"])
            for wn in warns:
                w.writerow([wn.get("type", ""), wn.get("text", ""), wn.get("basis", "")])
        return csv_path

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="594B3B")
    head_font = Font(bold=True, color="FFFFFF")
    tbd_fill = PatternFill("solid", fgColor="F6DAD2")

    ws = wb.active
    ws.title = "부서별 실행"
    ws.append(GRID_HEADER)
    for c in ws[1]:
        c.fill = head_fill; c.font = head_font; c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
        row = ws[ws.max_row]
        # 담당/자원 미정 강조
        if "TBD" in str(r[4]) or "미정" in str(r[5]) or r[8]:
            for c in row:
                c.fill = tbd_fill
    for col, wd in zip("ABCDEFGHI", [10, 24, 26, 18, 14, 20, 22, 8, 16]):
        ws.column_dimensions[col].width = wd
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    if gaps or warns:
        ws2 = wb.create_sheet("공백·경고")
        ws2.append(["구분", "유형", "내용", "근거/행동"])
        for c in ws2[1]:
            c.fill = head_fill; c.font = head_font
        for g in gaps:
            ws2.append(["공백", g.get("type", ""), g.get("text", ""), str(g.get("action_id", ""))])
        for wn in warns:
            ws2.append(["경고", wn.get("type", ""), wn.get("text", ""), wn.get("basis", "")])
        for col, wd in zip("ABCD", [10, 14, 40, 24]):
            ws2.column_dimensions[col].width = wd

    wb.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 2:
        print("사용: python3 build_action.py <action.json> [출력디렉토리]"); sys.exit(2)
    src = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(src))
    try:
        with open(src, encoding="utf-8") as f:
            a = json.load(f)
    except FileNotFoundError:
        print(f"[오류] 파일 없음: {src}"); sys.exit(2)
    except json.JSONDecodeError as ex:
        print(f"[오류] JSON 파싱 실패: {ex}"); sys.exit(2)

    if not isinstance(a, dict):
        print(f"[오류] 최상위 JSON이 객체(dict)가 아님: {type(a).__name__} — action.json 구조 확인"); sys.exit(2)

    errs = gate(a)
    if errs:
        print("=" * 60)
        print("렌더 거부 — 게이트 위반 (fidelity-guard 먼저 통과시키라)")
        print("=" * 60)
        for er in errs:
            print(f"  ✗ {er}")
        sys.exit(1)

    os.makedirs(out_dir, exist_ok=True)
    md_path = os.path.join(out_dir, "action-plan.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(build_md(a))
    xlsx_path = build_xlsx(a, os.path.join(out_dir, "action-plan.xlsx"))

    print("렌더 완료 ✓")
    print(f"  · {md_path}")
    print(f"  · {xlsx_path}" + ("  (openpyxl 없어 csv 폴백)" if xlsx_path.endswith(".csv") else ""))
    sys.exit(0)


if __name__ == "__main__":
    main()
