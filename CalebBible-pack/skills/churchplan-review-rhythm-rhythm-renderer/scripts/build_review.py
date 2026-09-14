#!/usr/bin/env python3
"""
build_review.py — churchplan-review-rhythm 점검 리듬 렌더러.

review.json → ① review-scorecard.html (분기 신호등 카드 4분기 + 월간 5분 점검 + 빨강 대응 + 연간 일정·반기 수정일 강조)
              ② review-rhythm.xlsx (분기 스코어카드 시트 + 월간 점검표 시트 + 점검 일정 시트, openpyxl 없으면 .csv)

렌더 전 핵심 게이트를 재강제한다(cys-insight 패턴):
  필수 블록 · 표어 · 미래값 둔갑 금지 · 반기 수정일 · status. 위반 시 렌더 거부(종료코드 1).
전체 정밀 검증(돌봄 프레이밍·상속 충실)은 fidelity-guard의 validate_review.py가 맡는다.

사용:  python3 build_review.py <review.json> [출력디렉토리]
종료코드: 0 = 성공, 1 = 게이트 위반(렌더 거부), 2 = 파일/JSON 오류
표준 라이브러리만 사용(xlsx는 openpyxl 있을 때만, 없으면 csv).
"""
import csv
import html
import json
import os
import sys


def _d(x):
    """truthy 비-dict(문자열·리스트·정수)도 {}로 안전 환원."""
    return x if isinstance(x, dict) else {}


def _l(x):
    """truthy 비-list도 []로 안전 환원."""
    return x if isinstance(x, list) else []


def _ld(x):
    """list-of-dict만 남긴다(비-dict 원소 제거 — 산출물 오염 차단)."""
    return [v for v in x if isinstance(v, dict)] if isinstance(x, list) else []


def _cell(x):
    """xlsx/csv 셀 안전 값. dict·list는 산출물 오염을 막으려 빈칸으로,
    그 외(str·int·float·bool·None)는 그대로 둔다(openpyxl이 처리)."""
    if isinstance(x, (dict, list)):
        return ""
    return x


def esc_plain(x):
    """f-string 내장용 안전 문자열. dict·list는 빈칸, None은 빈칸, 그 외는 str()."""
    if isinstance(x, (dict, list)) or x is None:
        return ""
    return str(x)


def gate(r):
    """렌더 전 핵심 게이트. truthy 비-dict/비-list 오염도 raw traceback 없이 위반 등록."""
    errs = []
    for blk in ("meta", "scorecard", "schedule"):
        if blk not in r:
            errs.append(f"필수 블록 '{blk}' 없음")
        elif not isinstance(r.get(blk), dict):
            errs.append(f"'{blk}'가 객체(dict)가 아님: {type(r.get(blk)).__name__}")
    meta = _d(r.get("meta"))
    theme = meta.get("chosen_theme")
    if not theme or not str(theme).strip():
        errs.append("meta.chosen_theme 비어 있음 — 표어 먼저 확정")
    if meta.get("status") not in ("제안", "확정"):
        errs.append("meta.status 가 제안|확정 이 아님")
    # 미래값 둔갑: confirmed 전인데 status 채워짐
    if not r.get("confirmed"):
        sc = _d(r.get("scorecard"))
        quarters_raw = sc.get("quarters")
        if "quarters" in sc and not isinstance(quarters_raw, list):
            errs.append(f"scorecard.quarters가 배열(list)이 아님: {type(quarters_raw).__name__}")
        for q in _l(quarters_raw):
            if not isinstance(q, dict):
                errs.append(f"scorecard.quarters 항목이 객체(dict)가 아님: {type(q).__name__}")
                continue
            for ind in _l(q.get("indicators")):
                if not isinstance(ind, dict):
                    errs.append(f"{q.get('quarter','?')} indicators 항목이 객체(dict)가 아님: {type(ind).__name__}")
                    continue
                if ind.get("status") not in (None, "", "null"):
                    errs.append(f"{q.get('quarter','?')} '{ind.get('ref','?')}' status가 점검 전에 채워짐 — 미래값 둔갑 금지")
    # 반기 수정일
    sch = _d(r.get("schedule"))
    mid = _d(sch.get("midyear_revision"))
    if mid.get("official") is not True:
        errs.append("schedule.midyear_revision(official:true) 없음 — 반기 공식 수정일 못박기")

    # 구조 타입 검사 — 알려진 dict/list 하위 필드가 present-but-wrong-type면 렌더 거부.
    # 검증기(validate_review.py)와 parity. 안전 환원(_d/_l/_ld)이 raw traceback은 막되,
    # 오염 입력은 산출물을 내지 않고 EXIT 1로 깔끔히 거부한다(과잉 차단 아님 — present-but-wrong-type만).
    def _nd(container, key, label):
        if isinstance(container, dict) and key in container and not isinstance(container[key], dict):
            errs.append(f"{label}가 객체(dict)가 아님: {type(container[key]).__name__}")

    def _nl(container, key, label):
        if isinstance(container, dict) and key in container and not isinstance(container[key], list):
            errs.append(f"{label}가 배열(list)이 아님: {type(container[key]).__name__}")

    if "inputs_echo" in r and not isinstance(r.get("inputs_echo"), dict):
        errs.append(f"inputs_echo가 객체(dict)가 아님: {type(r.get('inputs_echo')).__name__}")
    ie = _d(r.get("inputs_echo"))
    _nl(ie, "indicators_from_goals", "inputs_echo.indicators_from_goals")
    for i, el in enumerate(_l(ie.get("indicators_from_goals"))):
        if not isinstance(el, dict):
            errs.append(f"inputs_echo.indicators_from_goals[{i}] 항목이 객체(dict)가 아님: {type(el).__name__}")
    mc_blk = r.get("monthly_checkup")
    if "monthly_checkup" in r and not isinstance(mc_blk, dict):
        errs.append(f"monthly_checkup가 객체(dict)가 아님: {type(mc_blk).__name__}")
    if "red_signal_scenarios" in r and not isinstance(r.get("red_signal_scenarios"), list):
        errs.append(f"red_signal_scenarios가 배열(list)이 아님: {type(r.get('red_signal_scenarios')).__name__}")
    _nd(_d(mc_blk), "pastor_sustainability_row", "monthly_checkup.pastor_sustainability_row")
    _nl(_d(mc_blk), "rows", "monthly_checkup.rows")
    _nd(sch, "auto_reminder", "schedule.auto_reminder")
    _nl(sch, "quarterly_reviews", "schedule.quarterly_reviews")
    _nl(sch, "quarterly_agenda", "schedule.quarterly_agenda")
    for q in _ld(_d(r.get("scorecard")).get("quarters")):
        _nl(q, "indicators", f"{q.get('quarter','?')}.indicators")
    for i, el in enumerate(_l(r.get("red_signal_scenarios"))):
        if not isinstance(el, dict):
            errs.append(f"red_signal_scenarios[{i}] 항목이 객체(dict)가 아님: {type(el).__name__}")
    for i, el in enumerate(_l(_d(mc_blk).get("rows"))):
        if not isinstance(el, dict):
            errs.append(f"monthly_checkup.rows[{i}] 항목이 객체(dict)가 아님: {type(el).__name__}")
    for i, el in enumerate(_l(sch.get("quarterly_reviews"))):
        if not isinstance(el, dict):
            errs.append(f"schedule.quarterly_reviews[{i}] 항목이 객체(dict)가 아님: {type(el).__name__}")
    return errs


def esc(x):
    # dict·list 원소는 산출물 오염을 막으려 빈칸 처리(구조값을 HTML에 흘리지 않음).
    if isinstance(x, (dict, list)) or x is None:
        x = ""
    return html.escape(str(x))


def build_html(r):
    meta = _d(r.get("meta"))
    sc = _d(r.get("scorecard"))
    mc = _d(r.get("monthly_checkup"))
    sch = _d(r.get("schedule"))
    theme = meta.get("chosen_theme", "(미확정)")
    confirmed = bool(r.get("confirmed"))

    P = []
    P.append("<!DOCTYPE html><html lang='ko'><head><meta charset='utf-8'>")
    P.append("<meta name='viewport' content='width=device-width,initial-scale=1'>")
    P.append(f"<title>2027 점검 리듬 — {esc(theme)}</title>")
    P.append("""<style>
:root{--sepia:#594B3B;--ink:#2c2622;--paper:#FBF8F3;--card:#fff;--line:#E7DFD3;
--green:#5B8C5A;--yellow:#D9A441;--red:#C75A4A;--soft:#F3EDE3;--protect:#EAF0F7;}
*{box-sizing:border-box}body{margin:0;background:var(--paper);color:var(--ink);
font-family:'Nanum Myeongjo','Apple SD Gothic Neo',serif;line-height:1.6;padding:28px}
.wrap{max-width:920px;margin:0 auto}
h1{color:var(--sepia);font-size:25px;border-bottom:2px solid var(--sepia);padding-bottom:10px;margin:0 0 4px}
.theme{color:var(--sepia);font-size:15px;margin:0 0 2px}
.sub{color:#8a7f70;font-size:12.5px;margin:0 0 20px}
h2{color:var(--sepia);font-size:17px;margin:26px 0 10px;border-left:4px solid var(--sepia);padding-left:9px}
.frame{background:var(--soft);border-radius:9px;padding:11px 14px;font-size:13px;color:#5e5446;margin-bottom:14px}
.qgrid{display:grid;grid-template-columns:repeat(2,1fr);gap:13px}
.qcard{background:var(--card);border:1px solid var(--line);border-radius:11px;padding:13px 15px}
.qcard h3{margin:0 0 3px;color:var(--sepia);font-size:15px}
.qwhen{color:#9a8f80;font-size:11.5px;margin:0 0 9px}
.ind{border-top:1px dashed var(--line);padding:8px 0}
.ind:first-of-type{border-top:none}
.ind .lab{font-size:13.5px;font-weight:bold}
.ind .meta{font-size:11.5px;color:#8a7f70;margin:1px 0}
.ind .care{font-size:11.5px;color:#6f6354;font-style:italic;margin-top:2px}
.lights{display:inline-flex;gap:5px;vertical-align:middle;margin-left:6px}
.dot{width:13px;height:13px;border-radius:50%;border:2px solid;display:inline-block}
.dg{border-color:var(--green)}.dy{border-color:var(--yellow)}.dr{border-color:var(--red)}
.fill .dg{background:var(--green)}.fill .dy{background:var(--yellow)}.fill .dr{background:var(--red)}
.verify{display:inline-block;background:#FBEFD6;color:#8a6d1f;border-radius:5px;font-size:10.5px;padding:1px 6px;margin-left:5px}
.note{font-size:11px;color:#9a8f80;margin-top:6px}
table{border-collapse:collapse;width:100%;font-size:13px;margin-top:6px}
th,td{border:1px solid var(--line);padding:7px 9px;text-align:left;vertical-align:top}
th{background:var(--sepia);color:#fff;font-weight:600}
tr.protect td{background:var(--protect)}
.red-card{background:#FCF3F1;border:1px solid #EBC9C1;border-radius:9px;padding:11px 14px;margin-bottom:10px}
.red-card .h{color:var(--red);font-weight:bold;font-size:13.5px}
.red-card ul{margin:6px 0 4px;padding-left:18px;font-size:12.5px}
.red-card .not{font-size:11.5px;color:#9a6258}
.mid{background:#EFF4EF;border:1px solid #CFE0CE;border-radius:9px;padding:11px 14px;margin-top:8px}
.mid .h{color:var(--green);font-weight:bold}
.badge{display:inline-block;background:var(--sepia);color:#fff;border-radius:5px;font-size:11px;padding:1px 8px}
.footer{margin-top:26px;border-top:1px solid var(--line);padding-top:10px;color:#9a8f80;font-size:11.5px;text-align:center}
</style></head><body><div class='wrap'>""")

    P.append(f"<h1>2027 점검 리듬 스코어카드</h1>")
    P.append(f"<p class='theme'>표어 · {esc(theme)}</p>")
    P.append(f"<p class='sub'>상태: {'✅ 목사 확정' if confirmed else '🙏 목사 확정 대기'} &nbsp;·&nbsp; 신호등 색은 점검 시점에 채웁니다(미래값 비움)</p>")

    # 신호등 frame
    if sc.get("frame"):
        P.append(f"<div class='frame'>🚦 {esc(sc.get('frame'))}</div>")

    # 분기 카드
    P.append("<h2>분기 신호등 (1Q~4Q)</h2>")
    P.append("<div class='qgrid'>")
    for q in _ld(sc.get("quarters")):
        P.append("<div class='qcard'>")
        P.append(f"<h3>{esc(q.get('quarter',''))}</h3>")
        P.append(f"<p class='qwhen'>점검: {esc(q.get('review_when','분기 말'))}</p>")
        inds = _ld(q.get("indicators"))
        if not inds:
            P.append("<p class='note'>이 분기 점검 지표 없음</p>")
        for ind in inds:
            filled = ind.get("status") not in (None, "", "null")
            fillcls = " fill" if filled else ""
            verify = "<span class='verify'>실측 먼저</span>" if ind.get("verify_at_review") else ""
            P.append("<div class='ind'>")
            P.append(f"<div class='lab'>{esc(ind.get('ref',''))}{verify}"
                     f"<span class='lights{fillcls}'><span class='dot dg'></span>"
                     f"<span class='dot dy'></span><span class='dot dr'></span></span></div>")
            P.append(f"<div class='meta'>기준선 {esc(ind.get('baseline',''))} → 올해 {esc(ind.get('target',''))} ({esc(ind.get('by',''))})</div>")
            if ind.get("care_note"):
                P.append(f"<div class='care'>“{esc(ind.get('care_note'))}”</div>")
            P.append("</div>")
        P.append("</div>")
    P.append("</div>")

    # 월간 점검표
    if mc:
        P.append("<h2>월간 1페이지 점검 (5분)</h2>")
        if mc.get("frame"):
            P.append(f"<div class='frame'>🗓️ {esc(mc.get('frame'))}</div>")
        P.append("<table><tr><th>항목</th><th>출처</th><th>한눈 질문(5분)</th></tr>")
        for row in _ld(mc.get("rows")):
            P.append(f"<tr><td>{esc(row.get('item',''))}</td><td>{esc(row.get('source',''))}</td>"
                     f"<td>{esc(row.get('glance_question',''))}</td></tr>")
        psr = _d(mc.get("pastor_sustainability_row"))
        if psr:
            P.append(f"<tr class='protect'><td>🛡️ {esc(psr.get('item',''))}</td><td>{esc(psr.get('source',''))}</td>"
                     f"<td>{esc(psr.get('glance_question',''))}</td></tr>")
        P.append("</table>")
        if mc.get("weekly_option"):
            P.append(f"<p class='note'>주간 옵션: {esc(mc.get('weekly_option'))}</p>")

    # 빨강 대응
    scen = _ld(r.get("red_signal_scenarios"))
    if scen:
        P.append("<h2>🔴 빨강 지표 돌봄 대응</h2>")
        for sc1 in scen:
            P.append("<div class='red-card'>")
            P.append(f"<div class='h'>{esc(sc1.get('indicator_ref',''))}</div>")
            if sc1.get("if_red"):
                P.append(f"<div>{esc(sc1.get('if_red'))}</div>")
            resp = _l(sc1.get("response"))
            if resp:
                P.append("<ul>" + "".join(f"<li>{esc(x)}</li>" for x in resp if not isinstance(x, (dict, list))) + "</ul>")
            if sc1.get("not"):
                P.append(f"<div class='not'>※ {esc(sc1.get('not'))}</div>")
            P.append("</div>")

    # 연간 일정
    P.append("<h2>연간 점검 일정</h2>")
    qrs = _ld(sch.get("quarterly_reviews"))
    if qrs:
        P.append("<table><tr><th>분기</th><th>시점</th></tr>")
        for qr in qrs:
            P.append(f"<tr><td>{esc(qr.get('quarter',''))}</td><td>{esc(qr.get('when',''))}</td></tr>")
        P.append("</table>")
    agenda = [a for a in _l(sch.get("quarterly_agenda")) if not isinstance(a, (dict, list))]
    if agenda:
        P.append("<p class='note'>분기 어젠다: " + " · ".join(esc(a) for a in agenda) + "</p>")
    mid = _d(sch.get("midyear_revision"))
    if mid:
        P.append(f"<div class='mid'><span class='h'>📌 반기 공식 수정일</span> "
                 f"&nbsp;<span class='badge'>{esc(mid.get('when',''))}</span><br>{esc(mid.get('note',''))}</div>")
    ar = _d(sch.get("auto_reminder"))
    if ar and ar.get("proposed"):
        P.append(f"<p class='note'>(선택) 자동 리마인드 구조 제안: {esc(ar.get('structure',''))} — {esc(ar.get('note',''))}</p>")

    P.append("<div class='footer'>숫자는 목자의 눈이지 채찍이 아닙니다. 점검은 돌봄이 닿지 않은 양을 찾는 시간입니다.<br>"
             "신호등 색·진행값은 점검 시점에 채웁니다 — 계획은 갱신으로 살아 움직입니다.</div>")
    P.append("</div></body></html>")
    return "\n".join(P)


def build_xlsx(r, out_path):
    sc = _d(r.get("scorecard"))
    mc = _d(r.get("monthly_checkup"))
    sch = _d(r.get("schedule"))

    score_rows = []
    for q in _ld(sc.get("quarters")):
        for ind in _ld(q.get("indicators")):
            score_rows.append([
                _cell(q.get("quarter", "")), _cell(ind.get("ref", "")), _cell(ind.get("baseline", "")),
                _cell(ind.get("target", "")), _cell(ind.get("by", "")),
                "실측 먼저" if ind.get("verify_at_review") else "",
                _cell(ind.get("status")) or "(점검 시점에)",
            ])
    month_rows = [[_cell(row.get("item", "")), _cell(row.get("source", "")), _cell(row.get("glance_question", ""))]
                  for row in _ld(mc.get("rows"))]
    psr = _d(mc.get("pastor_sustainability_row"))
    if psr:
        month_rows.append([f"🛡️ {esc_plain(psr.get('item',''))}", _cell(psr.get("source", "")), _cell(psr.get("glance_question", ""))])
    sched_rows = [[_cell(qr.get("quarter", "")), _cell(qr.get("when", ""))] for qr in _ld(sch.get("quarterly_reviews"))]
    mid = _d(sch.get("midyear_revision"))
    if mid:
        sched_rows.append(["반기 공식 수정일", f"{esc_plain(mid.get('when',''))} — {esc_plain(mid.get('note',''))}"])

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        csv_path = out_path.rsplit(".", 1)[0] + ".csv"
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["[분기 스코어카드] 분기", "지표", "기준선", "올해 목표치", "기한", "실측", "점검 결과"])
            for row in score_rows:
                w.writerow(row)
            w.writerow([]); w.writerow(["[월간 점검표] 항목", "출처", "한눈 질문"])
            for row in month_rows:
                w.writerow(row)
            w.writerow([]); w.writerow(["[점검 일정] 시점", "내용"])
            for row in sched_rows:
                w.writerow(row)
        return csv_path

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="594B3B")
    head_font = Font(bold=True, color="FFFFFF")
    prot_fill = PatternFill("solid", fgColor="EAF0F7")

    ws = wb.active
    ws.title = "분기 스코어카드"
    ws.append(["분기", "지표", "기준선", "올해 목표치", "기한", "실측", "점검 결과"])
    for c in ws[1]:
        c.fill = head_fill; c.font = head_font
    for row in score_rows:
        ws.append(row)
    for col, wd in zip("ABCDEFG", [7, 30, 18, 18, 7, 10, 14]):
        ws.column_dimensions[col].width = wd

    ws2 = wb.create_sheet("월간 점검표")
    ws2.append(["항목", "출처", "한눈 질문(5분)"])
    for c in ws2[1]:
        c.fill = head_fill; c.font = head_font
    for row in month_rows:
        ws2.append(row)
        if str(row[0]).startswith("🛡️"):
            for c in ws2[ws2.max_row]:
                c.fill = prot_fill
    for col, wd in zip("ABC", [22, 22, 40]):
        ws2.column_dimensions[col].width = wd

    ws3 = wb.create_sheet("점검 일정")
    ws3.append(["시점", "내용"])
    for c in ws3[1]:
        c.fill = head_fill; c.font = head_font
    for row in sched_rows:
        ws3.append(row)
    for col, wd in zip("AB", [18, 52]):
        ws3.column_dimensions[col].width = wd

    wb.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 2:
        print("사용: python3 build_review.py <review.json> [출력디렉토리]"); sys.exit(2)
    src = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(src))
    try:
        with open(src, encoding="utf-8") as f:
            r = json.load(f)
    except FileNotFoundError:
        print(f"[오류] 파일 없음: {src}"); sys.exit(2)
    except json.JSONDecodeError as ex:
        print(f"[오류] JSON 파싱 실패: {ex}"); sys.exit(2)

    if not isinstance(r, dict):
        print(f"[오류] 최상위가 객체(dict)가 아님: {type(r).__name__} — review.json은 객체여야 함"); sys.exit(2)

    errs = gate(r)
    if errs:
        print("=" * 60)
        print("렌더 거부 — 게이트 위반 (fidelity-guard 먼저 통과시키라)")
        print("=" * 60)
        for er in errs:
            print(f"  ✗ {er}")
        sys.exit(1)

    os.makedirs(out_dir, exist_ok=True)
    html_path = os.path.join(out_dir, "review-scorecard.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(build_html(r))
    xlsx_path = build_xlsx(r, os.path.join(out_dir, "review-rhythm.xlsx"))

    print("렌더 완료 ✓")
    print(f"  · {html_path}")
    print(f"  · {xlsx_path}" + ("  (openpyxl 없어 csv 폴백)" if xlsx_path.endswith(".csv") else ""))
    sys.exit(0)


if __name__ == "__main__":
    main()
