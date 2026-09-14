#!/usr/bin/env python3
"""
build_package.py — churchplan-plan-compiler 통합 패키지 렌더러(자기완결).

plan-package.json → ① plan-document.docx (계획서 본문 9요소·나눔명조)  ② plan-document.pdf (weasyprint)
                    ③ slide-deck.pptx (당회/제직회 10장)
                    ④ vision-onepager.html + .pdf (성도용·은혜 톤)
                    ⑤ reinforce-guide.md (옵션2 보강 안내·#12·#9 다리)
                    ⑥ bundle.md (앞 산출물 참조 목록 — 재생산 아님)

코어 출력 스킬(docx·pptx·pdf)에 의존하지 않는다 — python-docx·python-pptx·weasyprint 자기완결.
라이브러리 없으면 폴백(docx→md·pptx→md·pdf 생략).

렌더 전 핵심 게이트를 재강제한다(cys-insight 패턴):
  필수 블록 · 표어 · 확정 게이트(missing/unconfirmed) · 9요소 · status. 위반 시 렌더 거부(종료코드 1).
전체 정밀 검증(전사 충실·할루시네이션)은 fidelity-guard의 validate_package.py가 맡는다.

사용:  python3 build_package.py <plan-package.json> [출력디렉토리]
종료코드: 0 = 성공, 1 = 게이트 위반(렌더 거부), 2 = 파일/JSON 오류
"""
import html
import json
import os
import sys

NINE = ["표어", "방향 선언문", "3축 목표", "설교계획", "핵심행사",
        "실행계획", "예산", "조직", "점검계획"]
FONT = "나눔명조"
SEPIA = "594B3B"


def _d(x):
    return x if isinstance(x, dict) else {}


def _l(x):
    return x if isinstance(x, list) else []


def _ld(x):
    return [v for v in x if isinstance(v, dict)] if isinstance(x, list) else []


def gate(p):
    """렌더 전 핵심 게이트."""
    errs = []
    for blk in ("meta", "sources", "plan_document", "slide_deck", "vision_onepager"):
        if blk not in p:
            errs.append(f"필수 블록 '{blk}' 없음")
    # 최상위 필수 dict 필드가 비-dict면 [형식오류] 위반 1건(조용한 통과 금지)
    for blk in ("meta", "sources", "plan_document", "slide_deck", "vision_onepager"):
        if blk in p and not isinstance(p.get(blk), dict):
            errs.append(f"최상위 '{blk}'가 객체(dict)가 아님: {type(p.get(blk)).__name__}")
    meta = _d(p.get("meta"))
    if not str(meta.get("chosen_theme") or "").strip():
        errs.append("meta.chosen_theme 비어 있음 — 표어 먼저 확정")
    if meta.get("status") not in ("제안", "확정"):
        errs.append("meta.status 가 제안|확정 이 아님")
    sources = _d(p.get("sources"))
    if _l(sources.get("missing")):
        errs.append(f"sources.missing 비어 있지 않음: {_l(sources.get('missing'))} — 빠진 산출물 먼저")
    if _l(sources.get("unconfirmed")):
        errs.append(f"sources.unconfirmed 비어 있지 않음: {_l(sources.get('unconfirmed'))} — 미확정 산출물 문서화 금지")
    chapters_raw = _d(p.get("plan_document")).get("chapters")
    if "chapters" in _d(p.get("plan_document")) and not isinstance(chapters_raw, list):
        errs.append("plan_document.chapters가 리스트가 아님")
    chapters = _l(chapters_raw)
    if any(not isinstance(c, dict) for c in chapters):
        errs.append("plan_document.chapters에 객체(dict) 아닌 항목 있음 — 각 장은 객체여야 함")
    keys = [c.get("key") for c in chapters if isinstance(c, dict)]
    for need in NINE:
        if need not in keys:
            errs.append(f"plan_document.chapters에 '{need}' 없음 — 9요소 고정")
    # 확정 게이트 — status=확정인데 confirmed 없으면 렌더 거부(목사 HITL 전 확정 표기 금지)
    if meta.get("status") == "확정" and not p.get("confirmed"):
        errs.append("status='확정'인데 confirmed 없음 — 목사 검토·확정 HITL 전 렌더 금지")

    # 타입 위반 게이트(검증기 parity) — 필드가 *있는데* 컨테이너 타입이 틀리면 렌더 거부.
    # 정상 입력(필드 없음 또는 올바른 타입)은 회귀 없이 통과(과잉 차단 금지).
    def _tv(container, key, want, label):
        if isinstance(container, dict) and key in container:
            v = container.get(key)
            if want == "dict" and not isinstance(v, dict):
                errs.append(f"{label}이(가) 객체(dict)가 아님: {type(v).__name__}")
            elif want == "list" and not isinstance(v, list):
                errs.append(f"{label}이(가) 리스트가 아님: {type(v).__name__}")

    def _bad_scalar(container, key, label):
        if isinstance(container, dict) and isinstance(container.get(key), list):
            bad = [i for i, v in enumerate(container[key]) if isinstance(v, (dict, list))]
            if bad:
                errs.append(f"{label}에 dict/list 항목 {bad} — 문자열 등 단순값이어야 함")

    pd_raw = p.get("plan_document")
    deck_raw = p.get("slide_deck")
    vp_raw = p.get("vision_onepager")
    rg_raw = p.get("reinforce_guide")
    _tv(_d(sources), "missing", "list", "sources.missing")
    _tv(_d(sources), "unconfirmed", "list", "sources.unconfirmed")
    _tv(pd_raw, "intro_optional", "dict", "plan_document.intro_optional")
    _tv(deck_raw, "slides", "list", "slide_deck.slides")
    _tv(vp_raw, "directions", "list", "vision_onepager.directions")
    if "reinforce_guide" in p and not isinstance(rg_raw, dict):
        errs.append(f"reinforce_guide가 객체(dict)가 아님: {type(rg_raw).__name__}")
    _tv(rg_raw, "verify_seeds", "list", "reinforce_guide.verify_seeds")
    if "bundle" in p and not isinstance(p.get("bundle"), list):
        errs.append(f"bundle이 리스트가 아님: {type(p.get('bundle')).__name__}")
    # 리스트 원소 타입 위반
    if isinstance(deck_raw, dict) and isinstance(deck_raw.get("slides"), list):
        bad = [i for i, v in enumerate(deck_raw["slides"]) if not isinstance(v, dict)]
        if bad:
            errs.append(f"slide_deck.slides에 dict 아닌 항목 {bad}")
    if isinstance(rg_raw, dict) and isinstance(rg_raw.get("verify_seeds"), list):
        bad = [i for i, v in enumerate(rg_raw["verify_seeds"]) if not isinstance(v, dict)]
        if bad:
            errs.append(f"reinforce_guide.verify_seeds에 dict 아닌 항목 {bad}")
    if isinstance(p.get("bundle"), list):
        bad = [i for i, v in enumerate(p["bundle"]) if not isinstance(v, dict)]
        if bad:
            errs.append(f"bundle에 dict 아닌 항목 {bad}")
    _bad_scalar(vp_raw, "directions", "vision_onepager.directions")
    for _i, _sl in enumerate(_l(deck_raw.get("slides") if isinstance(deck_raw, dict) else None)):
        _tv(_sl, "bullets", "list", f"slide_deck.slides[{_i}].bullets")
        _bad_scalar(_sl, "bullets", f"slide_deck.slides[{_i}].bullets")
    for _i, _bn in enumerate(_l(p.get("bundle"))):
        _tv(_bn, "files", "list", f"bundle[{_i}].files")
        _bad_scalar(_bn, "files", f"bundle[{_i}].files")
    return errs


def esc(x):
    return html.escape(str(x if x is not None else ""))


# ── 계획서 본문 ──
def plan_html(p):
    meta = _d(p.get("meta"))
    pd = _d(p.get("plan_document"))
    name = meta.get("church_name", "○○교회")
    theme = meta.get("chosen_theme", "")
    year = meta.get("year", 2027)
    L = ["<!DOCTYPE html><html lang='ko'><head><meta charset='utf-8'>",
         f"<title>{esc(name)} {esc(year)} 목회계획서</title>",
         f"""<style>
@page{{size:A4;margin:22mm 18mm;@bottom-center{{content:counter(page);color:#8a8a8a;font-size:9pt}}}}
body{{font-family:'{FONT}','Apple SD Gothic Neo',serif;color:#1f1b17;line-height:1.7;font-size:11pt}}
h1{{color:#{SEPIA};font-size:22pt;text-align:center;margin:0 0 2pt}}
.theme{{text-align:center;color:#{SEPIA};font-size:13pt;margin:0 0 4pt}}
.sub{{text-align:center;color:#8a7f70;font-size:10pt;margin:0 0 20pt}}
h2{{color:#{SEPIA};font-size:14pt;border-bottom:1.5pt solid #{SEPIA};padding-bottom:3pt;margin:18pt 0 7pt}}
.star{{color:#b5793a}}
.src{{color:#a99;font-size:8.5pt}}
p{{margin:4pt 0;text-align:justify}}
.intro{{background:#F6F1E8;border-radius:6pt;padding:8pt 11pt;font-size:10pt;color:#5e5446}}
</style></head><body>"""]
    L.append(f"<h1>{esc(name)} {esc(year)} 목회계획서</h1>")
    L.append(f"<p class='theme'>{esc(theme)}</p>")
    L.append(f"<p class='sub'>{esc(year)} · {esc(meta.get('tone',''))}</p>")
    intro = _d(pd.get("intro_optional"))
    if intro.get("content"):
        L.append(f"<div class='intro'><b>지난해 돌아봄</b><br>{esc(intro.get('content'))}</div>")
    for c in _ld(pd.get("chapters")):
        star = " <span class='star'>★</span>" if c.get("key") in ("설교계획", "핵심행사") else ""
        L.append(f"<h2>{esc(c.get('key',''))}{star}</h2>")
        for para in str(c.get("content", "")).split("\n"):
            if para.strip():
                L.append(f"<p>{esc(para)}</p>")
        L.append(f"<p class='src'>출처: {esc(c.get('source',''))}</p>")
    L.append("<p class='src' style='margin-top:18pt'>이 계획서는 목사의 분별·기도의 결과를 옮긴 것입니다. "
             "Claude는 문서화를 도왔을 뿐입니다.</p>")
    L.append("</body></html>")
    return "\n".join(L)


def plan_docx(p, path):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
    except ImportError:
        return None
    meta = _d(p.get("meta"))
    pd = _d(p.get("plan_document"))
    sepia = RGBColor(0x59, 0x4B, 0x3B)

    def setfont(run, size, color=None, bold=False):
        run.font.name = FONT
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        run.font.bold = bold
        rPr = run._element.get_or_add_rPr()
        rF = rPr.find(qn("w:rFonts"))
        if rF is None:
            rF = rPr.makeelement(qn("w:rFonts"), {}); rPr.append(rF)
        for a in ("w:eastAsia", "w:ascii", "w:hAnsi"):
            rF.set(qn(a), FONT)

    doc = Document()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    setfont(t.add_run(f"{meta.get('church_name','○○교회')} {meta.get('year',2027)} 목회계획서"), 20, sepia, True)
    th = doc.add_paragraph(); th.alignment = WD_ALIGN_PARAGRAPH.CENTER
    setfont(th.add_run(meta.get("chosen_theme", "")), 13, sepia)
    intro = _d(pd.get("intro_optional"))
    if intro.get("content"):
        h = doc.add_paragraph(); setfont(h.add_run("지난해 돌아봄"), 13, sepia, True)
        b = doc.add_paragraph(); setfont(b.add_run(str(intro.get("content"))), 11)
    for c in _ld(pd.get("chapters")):
        star = " ★" if c.get("key") in ("설교계획", "핵심행사") else ""
        h = doc.add_paragraph(); setfont(h.add_run(f"{c.get('key','')}{star}"), 14, sepia, True)
        for para in str(c.get("content", "")).split("\n"):
            if para.strip():
                bp = doc.add_paragraph(); setfont(bp.add_run(para), 11)
        s = doc.add_paragraph(); setfont(s.add_run(f"출처: {c.get('source','')}"), 8.5, RGBColor(0xA0, 0x99, 0x99))
    doc.save(path)
    return path


def plan_md(p):
    meta = _d(p.get("meta"))
    pd = _d(p.get("plan_document"))
    L = [f"# {meta.get('church_name','○○교회')} {meta.get('year',2027)} 목회계획서", "", f"**{meta.get('chosen_theme','')}**", ""]
    intro = _d(pd.get("intro_optional"))
    if intro.get("content"):
        L += ["## 지난해 돌아봄", "", str(intro.get("content")), ""]
    for c in _ld(pd.get("chapters")):
        star = " ★" if c.get("key") in ("설교계획", "핵심행사") else ""
        L += [f"## {c.get('key','')}{star}", "", str(c.get("content", "")), f"_출처: {c.get('source','')}_", ""]
    return "\n".join(L) + "\n"


# ── 슬라이드 ──
def build_pptx(p, path):
    try:
        from pptx import Presentation
        from pptx.util import Pt, Inches
        from pptx.dml.color import RGBColor
    except ImportError:
        return None
    meta = _d(p.get("meta"))
    deck = _d(p.get("slide_deck"))
    prs = Presentation()
    sepia = RGBColor(0x59, 0x4B, 0x3B)
    blank = prs.slide_layouts[6]
    for s in _ld(deck.get("slides")):
        slide = prs.slides.add_slide(blank)
        tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9), Inches(1))
        tf = tb.text_frame; tf.word_wrap = True
        r = tf.paragraphs[0].add_run(); r.text = f"{s.get('n','')}. {s.get('title','')}"
        r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = sepia; r.font.name = FONT
        bb = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(8.5), Inches(4.5))
        bf = bb.text_frame; bf.word_wrap = True
        bullets = [b for b in _l(s.get("bullets")) if not isinstance(b, (dict, list))]
        for i, bullet in enumerate(bullets):
            par = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
            run = par.add_run(); run.text = f"· {str(bullet)}"
            run.font.size = Pt(18); run.font.name = FONT
    prs.save(path)
    return path


def deck_md(p):
    deck = _d(p.get("slide_deck"))
    L = [f"# 당회/제직회 발표 슬라이드 ({deck.get('audience','')})", ""]
    for s in _ld(deck.get("slides")):
        L += [f"## {s.get('n','')}. {s.get('title','')}"]
        L += [f"- {str(b)}" for b in _l(s.get("bullets")) if not isinstance(b, (dict, list))]
        L += [f"_출처: {s.get('source','')}_", ""]
    return "\n".join(L) + "\n"


# ── 비전 한 장 ──
def vision_html(p):
    meta = _d(p.get("meta"))
    v = _d(p.get("vision_onepager"))
    name = meta.get("church_name", "○○교회")
    year = meta.get("year", 2027)
    L = ["<!DOCTYPE html><html lang='ko'><head><meta charset='utf-8'>",
         f"<title>{esc(name)} {esc(year)} 비전</title>",
         f"""<style>
@page{{size:A4;margin:24mm}}
body{{font-family:'{FONT}',serif;color:#2c2622;text-align:center;line-height:1.8}}
.name{{color:#8a7f70;font-size:13pt;margin-bottom:6pt}}
.theme{{color:#{SEPIA};font-size:26pt;font-weight:bold;margin:10pt 0}}
.scr{{color:#7a6f60;font-size:12pt;font-style:italic;margin-bottom:22pt}}
.dir{{text-align:left;max-width:480pt;margin:0 auto}}
.dir li{{margin:8pt 0;font-size:13pt}}
.inv{{margin-top:26pt;background:#F6F1E8;border-radius:9pt;padding:14pt 18pt;font-size:13pt;color:#5e5446}}
</style></head><body>"""]
    L.append(f"<p class='name'>{esc(name)} · {esc(year)}</p>")
    L.append(f"<p class='theme'>{esc(v.get('theme',''))}</p>")
    if v.get("core_scripture"):
        L.append(f"<p class='scr'>{esc(v.get('core_scripture'))}</p>")
    dirs = [d for d in _l(v.get("directions")) if not isinstance(d, (dict, list))]
    if dirs:
        L.append("<ul class='dir'>" + "".join(f"<li>{esc(d)}</li>" for d in dirs) + "</ul>")
    if v.get("invitation"):
        L.append(f"<div class='inv'>{esc(v.get('invitation'))}</div>")
    L.append("</body></html>")
    return "\n".join(L)


# ── 보강 안내 / 번들 ──
def reinforce_md(p):
    rg = _d(p.get("reinforce_guide"))
    L = ["# 옵션2 — 귀가 후 보강 안내 (계획은 살아 움직입니다)", "",
         "워크샵은 끝이 아니라 보강의 시작입니다. 아래는 귀가 후 실제 자료로 또렷이 할 항목입니다.", ""]
    seeds = _ld(rg.get("verify_seeds"))
    if seeds:
        L += ["## 귀가 후 실측 대상 (지금은 추정)", ""]
        for s in seeds:
            L.append(f"- **{s.get('indicator','')}** — {s.get('reason','')} _(출처: {s.get('from','')})_")
        L.append("")
    if rg.get("auto_review_note"):
        L += ["## 분기 자동 점검 (원하는 목사만)", "", str(rg.get("auto_review_note")), ""]
    if rg.get("next"):
        L += ["## 다음 다리", "", f"➡️ {rg.get('next')}", ""]
    return "\n".join(L) + "\n"


def bundle_md(p):
    L = ["# 번들 — 앞 단계 산출물 (재생산 아님·참조)", "",
         "아래 산출물은 앞 스킬이 이미 만든 파일입니다. 계획서 패키지에 함께 묶어 보관하세요.", ""]
    for b in _ld(p.get("bundle")):
        files = ", ".join(str(f) for f in _l(b.get("files")) if not isinstance(f, (dict, list)))
        L.append(f"- **{b.get('name','')}** ({b.get('from','')}) — {files}")
    return "\n".join(L) + "\n"


def render_pdf(html_str, path):
    try:
        from weasyprint import HTML
        HTML(string=html_str).write_pdf(path)
        return path
    except Exception:
        return None


def main():
    if len(sys.argv) < 2:
        print("사용: python3 build_package.py <plan-package.json> [출력디렉토리]"); sys.exit(2)
    src = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(src))
    try:
        with open(src, encoding="utf-8") as f:
            p = json.load(f)
    except FileNotFoundError:
        print(f"[오류] 파일 없음: {src}"); sys.exit(2)
    except json.JSONDecodeError as ex:
        print(f"[오류] JSON 파싱 실패: {ex}"); sys.exit(2)

    if not isinstance(p, dict):
        print(f"[오류] 최상위 JSON이 객체(dict)가 아님: {type(p).__name__} — plan-package.json은 객체여야 함"); sys.exit(2)

    errs = gate(p)
    if errs:
        print("=" * 60)
        print("렌더 거부 — 게이트 위반 (fidelity-guard 먼저 통과시키라)")
        print("=" * 60)
        for er in errs:
            print(f"  ✗ {er}")
        sys.exit(1)

    os.makedirs(out_dir, exist_ok=True)
    made = []

    # 계획서 (docx + pdf, 폴백 md)
    docx_path = plan_docx(p, os.path.join(out_dir, "plan-document.docx"))
    if docx_path:
        made.append(docx_path)
    else:
        mp = os.path.join(out_dir, "plan-document.md")
        with open(mp, "w", encoding="utf-8") as f:
            f.write(plan_md(p))
        made.append(mp + "  (python-docx 없어 md 폴백)")
    phtml = plan_html(p)
    pdf = render_pdf(phtml, os.path.join(out_dir, "plan-document.pdf"))
    made.append(pdf if pdf else "plan-document.pdf 생략(weasyprint 없음)")

    # 슬라이드 (pptx, 폴백 md)
    pptx_path = build_pptx(p, os.path.join(out_dir, "slide-deck.pptx"))
    if pptx_path:
        made.append(pptx_path)
    else:
        mp = os.path.join(out_dir, "slide-deck.md")
        with open(mp, "w", encoding="utf-8") as f:
            f.write(deck_md(p))
        made.append(mp + "  (python-pptx 없어 md 폴백)")

    # 비전 한 장 (html + pdf)
    vhtml = vision_html(p)
    vh = os.path.join(out_dir, "vision-onepager.html")
    with open(vh, "w", encoding="utf-8") as f:
        f.write(vhtml)
    made.append(vh)
    vpdf = render_pdf(vhtml, os.path.join(out_dir, "vision-onepager.pdf"))
    made.append(vpdf if vpdf else "vision-onepager.pdf 생략(weasyprint 없음)")

    # 보강 안내 + 번들
    rg = os.path.join(out_dir, "reinforce-guide.md")
    with open(rg, "w", encoding="utf-8") as f:
        f.write(reinforce_md(p))
    made.append(rg)
    bd = os.path.join(out_dir, "bundle.md")
    with open(bd, "w", encoding="utf-8") as f:
        f.write(bundle_md(p))
    made.append(bd)

    print("렌더 완료 ✓")
    for m in made:
        print(f"  · {m}")
    sys.exit(0)


if __name__ == "__main__":
    main()
