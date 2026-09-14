#!/usr/bin/env python3
"""
build_calendar.py — churchplan-ministry-calendar 통합 연간 사역 캘린더 렌더러.

calendar.json → ① calendar.html (분기·월·주 보기 + 부하 히트맵 + 안식 강조 + 밀집 경고)
                ② calendar.xlsx (52주 통합 그리드 시트 + 경고·안식 시트, openpyxl 없으면 .csv)

렌더 전 핵심 게이트를 재강제한다(cys-insight 패턴):
  필수 블록 · 표어 · 52주 무결성(범위·중복) · status. 위반 시 렌더 거부(종료코드 1).
전체 정밀 검증(상속 id 충실 등)은 fidelity-guard의 validate_calendar.py가 맡는다.

사용:  python3 build_calendar.py <calendar.json> [출력디렉토리]
종료코드: 0 = 성공, 1 = 게이트 위반(렌더 거부), 2 = 파일/JSON 오류
표준 라이브러리만 사용(xlsx는 openpyxl 있을 때만, 없으면 csv).
"""
import csv
import html
import json
import os
import sys

WEEKS_TOTAL = 52
SEPIA = "#594B3B"
LOAD_BG = {"낮음": "#E8F0E3", "보통": "#FBF3D9", "높음": "#F6DAD2"}
REST_BG = "#EAF0F7"
REST_BORDER = "#4A6FA5"


def _d(x):
    return x if isinstance(x, dict) else {}


def _l(x):
    return x if isinstance(x, list) else []


def _ld(x):
    return [v for v in x if isinstance(v, dict)] if isinstance(x, list) else []


def gate(c, preview=False):
    """렌더 전 핵심 게이트.

    preview=False(기본·최종 렌더): status=='확정' and confirmed 강제(목사 HITL 후에만).
    preview=True: 미확정 미리보기 허용(워터마크) — 최종 캘린더 아님.
    """
    errs = []
    for blk in ("meta", "weeks"):
        if blk not in c:
            errs.append(f"필수 블록 '{blk}' 없음")
    if "meta" in c and not isinstance(c.get("meta"), dict):
        errs.append("meta 가 객체(JSON object)가 아님 — 형식 오류")
    meta = _d(c.get("meta"))
    if not str(meta.get("chosen_theme") or "").strip():
        errs.append("meta.chosen_theme 비어 있음 — 표어 먼저 확정")
    status = meta.get("status")
    if status not in ("제안", "확정"):
        errs.append("meta.status 가 제안|확정 이 아님")
    # HITL 확정 게이트(최종 렌더 전용) — schema/SKILL: '확정 후에만 최종 캘린더 렌더'
    if not preview:
        if status != "확정":
            errs.append("meta.status 가 '확정'이 아님 — 목사 검토·확정(HITL) 후에만 최종 렌더(미확정은 --preview)")
        if not c.get("confirmed"):
            errs.append("confirmed 없음 — 목사 확정 HITL 전 최종 렌더 금지(미확정 미리보기는 --preview)")
    seen = set()
    for i, wk in enumerate(_l(c.get("weeks"))):
        if not isinstance(wk, dict):
            errs.append(f"weeks[{i}]: 객체여야 함"); continue
        n = wk.get("week")
        if not isinstance(n, int) or n < 1 or n > WEEKS_TOTAL:
            errs.append(f"week '{n}' 범위 이탈(1~{WEEKS_TOTAL})")
        elif n in seen:
            errs.append(f"week {n} 중복 — 격자 깨짐")
        else:
            seen.add(n)
    return errs


def esc(x):
    return html.escape(str(x)) if x is not None else ""


def week_cell(wk, density_idx):
    """한 주 셀 HTML."""
    n = wk.get("week")
    load = _d(wk.get("load"))
    level = load.get("level", "낮음")
    bg = LOAD_BG.get(level, "#F4F1EC")
    rest = wk.get("rest_protected")
    style = f"background:{REST_BG};border:2px solid {REST_BORDER};" if rest else f"background:{bg};"
    sm = _d(wk.get("sermon"))
    parts = [f'<div class="wk-no">W{n} · {esc(wk.get("date",""))}</div>']
    if wk.get("observance"):
        parts.append(f'<div class="obs">✝ {esc(wk.get("observance"))}</div>')
    if sm.get("title"):
        parts.append(f'<div class="sermon">📖 {esc(sm.get("title"))}</div>')
    for ev in _ld(wk.get("events")):
        et = ev.get("type")
        tag = "핵심" if et == "핵심" else ("정기" if et == "정기" else "미분류")
        parts.append(f'<div class="ev ev-{tag}">🎪 {esc(ev.get("title",""))}<span class="badge">{tag}</span></div>')
    for cp in _ld(wk.get("campaigns")):
        parts.append(f'<div class="camp">🔥 {esc(cp.get("name",""))}</div>')
    for ms in _ld(wk.get("milestones")):
        parts.append(f'<div class="ms">🎯 {esc(ms.get("label",""))} <small>({esc(ms.get("by",""))})</small></div>')
    if rest:
        parts.append('<div class="rest">🛌 안식 보호</div>')
    if n in density_idx:
        parts.append(f'<div class="dense">🔥 밀집 [{esc(density_idx[n])}]</div>')
    return f'<td class="wk" style="{style}">' + "".join(parts) + "</td>"


def build_html(c, preview=False):
    meta = _d(c.get("meta"))
    sf = _d(meta.get("season_frame"))
    weeks = {wk.get("week"): wk for wk in _l(c.get("weeks")) if isinstance(wk, dict) and isinstance(wk.get("week"), int)}

    # 밀집 인덱스: week -> level
    density_idx = {}
    for d in _ld(c.get("density_warnings")):
        for w in _l(d.get("weeks")):
            if isinstance(w, int):
                density_idx[w] = d.get("level", "")

    # 3중 정합·시즌 정렬 인덱스
    tri = {t.get("quarter"): t for t in _ld(c.get("tri_consistency"))}
    sa = {s.get("quarter"): s for s in _ld(c.get("season_alignment"))}

    confirmed = c.get("confirmed")
    stat = "✅ 목사 확정" if confirmed else "🙏 목사 확정 대기"

    H = []
    H.append("<!DOCTYPE html><html lang='ko'><head><meta charset='utf-8'>")
    H.append("<title>2027 통합 연간 사역 캘린더</title>")
    H.append(f"""<style>
    body{{font-family:'Nanum Gothic','맑은 고딕',sans-serif;background:#FAF8F4;color:#2B2620;margin:0;padding:24px;}}
    .head{{border-top:6px solid {SEPIA};border-bottom:1px solid #D8CFC2;padding:14px 0 10px;margin-bottom:8px;}}
    .head h1{{color:{SEPIA};margin:0 0 4px;font-size:24px;}}
    .head .theme{{font-size:17px;font-weight:bold;color:#3A332A;}}
    .head .meta{{color:#7A6F5F;font-size:13px;margin-top:4px;}}
    .legend{{font-size:12px;color:#5A5142;margin:10px 0 18px;display:flex;gap:14px;flex-wrap:wrap;}}
    .legend span{{padding:3px 8px;border-radius:4px;}}
    .qsec{{margin-bottom:26px;}}
    .qhead{{background:{SEPIA};color:#fff;padding:8px 12px;border-radius:6px 6px 0 0;font-size:15px;font-weight:bold;}}
    .qmeta{{background:#F0EAE0;padding:7px 12px;font-size:12.5px;color:#4A4234;border-left:3px solid {SEPIA};}}
    .qmeta .gap{{color:#B5482E;font-weight:bold;}}
    table{{border-collapse:collapse;width:100%;table-layout:fixed;}}
    .mhead td{{background:#EFE9DF;color:{SEPIA};font-weight:bold;font-size:12px;padding:4px 6px;border:1px solid #D8CFC2;}}
    td.wk{{vertical-align:top;border:1px solid #D8CFC2;padding:6px;height:96px;font-size:11.5px;line-height:1.45;width:14.2%;}}
    .wk-no{{font-weight:bold;color:{SEPIA};font-size:11px;}}
    .obs{{color:#7A5C2E;font-weight:bold;}}
    .sermon{{color:#2B2620;}}
    .ev{{color:#1F5C3A;margin-top:2px;}}
    .camp{{color:#9A4A1E;}}
    .ms{{color:#3A4A7A;}}
    .rest{{color:{REST_BORDER};font-weight:bold;margin-top:2px;}}
    .dense{{color:#B5482E;font-weight:bold;}}
    .badge{{font-size:9px;background:#1F5C3A;color:#fff;border-radius:3px;padding:0 4px;margin-left:3px;}}
    .ev-정기 .badge{{background:#8A7A5A;}}
    .ev-미분류 .badge{{background:#B5482E;}}
    .warnsec h2{{color:{SEPIA};font-size:17px;border-bottom:1px solid #D8CFC2;padding-bottom:4px;}}
    .warnsec li{{margin:5px 0;font-size:13px;}}
    .lv-고{{color:#B5482E;font-weight:bold;}} .lv-중{{color:#9A6A1E;}} .lv-저{{color:#5A5142;}}
    .note{{color:#7A6F5F;font-size:12px;font-style:italic;margin-top:18px;}}
    </style></head><body>""")

    if preview:
        H.append("<div style='background:#FBF3D9;border:2px solid #B5482E;color:#B5482E;"
                 "font-weight:bold;text-align:center;padding:8px;margin-bottom:10px;'>"
                 "⚠ 미리보기 — 목사 검토·확정(HITL) 전 초안입니다. 최종 캘린더가 아닙니다.</div>")
    H.append("<div class='head'>")
    H.append("<h1>2027 통합 연간 사역 캘린더</h1>")
    H.append(f"<div class='theme'>표어 · {esc(meta.get('chosen_theme'))}</div>")
    H.append(f"<div class='meta'>{esc(meta.get('year'))}년 · {stat} · 설교·행사·캠페인·목표 마일스톤 통합</div>")
    H.append("</div>")

    H.append("<div class='legend'>"
             f"<span style='background:{LOAD_BG['낮음']}'>부하 낮음(여백)</span>"
             f"<span style='background:{LOAD_BG['보통']}'>보통</span>"
             f"<span style='background:{LOAD_BG['높음']}'>높음</span>"
             f"<span style='background:{REST_BG};border:2px solid {REST_BORDER}'>🛌 안식 보호</span>"
             "<span>📖설교 🎪행사 🔥캠페인 🎯마일스톤</span></div>")

    # 분기 → 월 → 주 (세 보기를 한 페이지 계층으로)
    q_ranges = {1: range(1, 14), 2: range(14, 27), 3: range(27, 40), 4: range(40, 53)}
    for q in (1, 2, 3, 4):
        H.append("<div class='qsec'>")
        H.append(f"<div class='qhead'>{q}분기 · {esc(sf.get(f'{q}Q',''))}</div>")
        # 분기 메타(시즌 정렬 + 3중 정합)
        qm = []
        if q in sa:
            al = "정렬 ✓" if sa[q].get("aligned") else "어긋남 ⚠"
            qm.append(f"시즌: {al} {esc(sa[q].get('note',''))}")
        if q in tri:
            t = tri[q]
            if t.get("aligned"):
                qm.append(f"3중 정합 ✓ (설교·행사·사역 한 방향)")
            else:
                qm.append(f"<span class='gap'>3중 정합 ⚠ {esc(t.get('gap',''))}</span>")
        if qm:
            H.append(f"<div class='qmeta'>{' · '.join(qm)}</div>")
        # 월 그룹 → 주 행
        H.append("<table>")
        cur_month = None
        row_open = False
        col = 0
        for wn in q_ranges[q]:
            wk = weeks.get(wn)
            mo = wk.get("month") if wk else None
            if mo != cur_month:
                if row_open:
                    H.append("</tr>"); row_open = False
                H.append(f"<tr class='mhead'><td colspan='7'>{esc(mo)}월</td></tr>")
                cur_month = mo
                col = 0
            if col == 0:
                H.append("<tr>"); row_open = True
            H.append(week_cell(wk, density_idx) if wk else "<td class='wk' style='background:#F4F1EC'><div class='wk-no'>—</div></td>")
            col += 1
            if col == 7:
                H.append("</tr>"); row_open = False; col = 0
        if row_open:
            H.append("</tr>")
        H.append("</table></div>")

    # 밀집 경고
    H.append("<div class='warnsec'>")
    dw = _ld(c.get("density_warnings"))
    if dw:
        H.append("<h2>🔥 밀집·연속 고강도 경고</h2><ul>")
        for d in dw:
            lv = d.get("level", "")
            H.append(f"<li class='lv-{esc(lv)}'>[{esc(d.get('type',''))}] {esc(d.get('text',''))} "
                     f"<small>(주 {esc(d.get('weeks',''))} · 근거: {esc(d.get('basis',''))})</small></li>")
        H.append("</ul>")
    # 안식 계획
    rp = _ld(c.get("rest_plan"))
    H.append("<h2>🛌 안식 주간 보호 (여백도 계획이다)</h2>")
    if rp:
        H.append("<ul>")
        for r in rp:
            H.append(f"<li>주 {esc(r.get('weeks',''))} — <b>{esc(r.get('type',''))}</b> ({esc(r.get('status',''))}) "
                     f"{esc(r.get('note',''))} <small>{esc(r.get('basis',''))}</small></li>")
        H.append("</ul>")
    else:
        H.append("<p class='lv-중'>⚠ 안식 계획이 비어 있습니다 — 목자의 지속가능성을 점검하세요(가드레일 ④).</p>")
    # 상속 부하
    iw = _d(c.get("inherited_warnings"))
    inh = _ld(iw.get("from_sermon")) + _ld(iw.get("from_events"))
    if inh:
        H.append("<h2>⚖️ 상속된 총 부하 (설교 #4 · 행사 #5)</h2><ul>")
        for w in inh:
            H.append(f"<li>[{esc(w.get('type',''))}] {esc(w.get('text',''))}</li>")
        H.append("</ul>")
    H.append("</div>")

    H.append("<p class='note'>달력을 꽉 채우는 것이 목표가 아닙니다. 통합은 한 해를 한 눈에 정렬해 강단·행사·사역이 한 방향을 향하게 하고, 목자가 소진되지 않도록 여백을 지키는 일입니다.</p>")
    H.append("</body></html>")
    return "\n".join(H)


GRID_HEADER = ["주", "날짜", "분기", "월", "절기", "설교", "행사", "캠페인", "마일스톤", "부하"]


def grid_rows(c):
    rows = []
    weeks = _ld(c.get("weeks"))
    for wk in sorted(weeks, key=lambda x: x.get("week", 0) if isinstance(x.get("week"), int) else 0):
        sm = _d(wk.get("sermon"))
        load = _d(wk.get("load"))
        evs = _ld(wk.get("events"))
        cps = _ld(wk.get("campaigns"))
        mss = _ld(wk.get("milestones"))
        rows.append([
            wk.get("week", ""), wk.get("date", ""), wk.get("quarter", ""), wk.get("month", ""),
            wk.get("observance", ""), sm.get("title", ""),
            "; ".join(f"{e.get('title','')}({e.get('type','')})" for e in evs),
            "; ".join(str(cp.get("name", "")) for cp in cps),
            "; ".join(f"{m.get('label','')}[{m.get('by','')}]" for m in mss),
            f"{load.get('level','')}({load.get('score','')})",
        ])
    return rows


def build_xlsx(c, out_path):
    rows = grid_rows(c)
    dw = _ld(c.get("density_warnings"))
    rp = _ld(c.get("rest_plan"))
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        csv_path = out_path.rsplit(".", 1)[0] + ".csv"
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(GRID_HEADER)
            for r in rows:
                w.writerow(r)
            w.writerow([])
            w.writerow(["[밀집 경고]"])
            for d in dw:
                w.writerow([d.get("type", ""), d.get("level", ""), d.get("text", ""), str(d.get("weeks", "")), d.get("basis", "")])
            w.writerow([])
            w.writerow(["[안식 계획]"])
            for r in rp:
                w.writerow([str(r.get("weeks", "")), r.get("type", ""), r.get("status", ""), r.get("note", ""), r.get("basis", "")])
        return csv_path

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="594B3B")
    head_font = Font(bold=True, color="FFFFFF")
    lvl_fill = {"낮음": PatternFill("solid", fgColor="E8F0E3"),
                "보통": PatternFill("solid", fgColor="FBF3D9"),
                "높음": PatternFill("solid", fgColor="F6DAD2")}
    rest_fill = PatternFill("solid", fgColor="EAF0F7")

    ws = wb.active
    ws.title = "52주 통합"
    ws.append(GRID_HEADER)
    for cc in ws[1]:
        cc.fill = head_fill; cc.font = head_font; cc.alignment = Alignment(horizontal="center", vertical="center")
    rest_weeks = set()
    for r in rp:
        for w in _l(r.get("weeks")):
            if isinstance(w, (int, str)):
                rest_weeks.add(w)
    for r in rows:
        ws.append(r)
        row = ws[ws.max_row]
        lvl = str(r[9]).split("(")[0]
        fill = rest_fill if r[0] in rest_weeks else lvl_fill.get(lvl)
        if fill:
            for cell in row:
                cell.fill = fill
    for col, wd in zip("ABCDEFGHIJ", [5, 12, 6, 5, 12, 26, 28, 18, 24, 12]):
        ws.column_dimensions[col].width = wd
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws2 = wb.create_sheet("경고·안식")
    ws2.append(["구분", "유형/수준", "내용", "주차", "근거/비고"])
    for cc in ws2[1]:
        cc.fill = head_fill; cc.font = head_font
    for d in dw:
        ws2.append(["밀집경고", f"{d.get('type','')}/{d.get('level','')}", d.get("text", ""), str(d.get("weeks", "")), d.get("basis", "")])
    for r in rp:
        ws2.append(["안식계획", f"{r.get('type','')}/{r.get('status','')}", r.get("note", ""), str(r.get("weeks", "")), r.get("basis", "")])
    for col, wd in zip("ABCDE", [10, 16, 36, 14, 28]):
        ws2.column_dimensions[col].width = wd

    wb.save(out_path)
    return out_path


def main():
    args = [a for a in sys.argv[1:] if a != "--preview"]
    preview = "--preview" in sys.argv[1:]
    if not args:
        print("사용: python3 build_calendar.py <calendar.json> [출력디렉토리] [--preview]"); sys.exit(2)
    src = args[0]
    out_dir = args[1] if len(args) > 1 else os.path.dirname(os.path.abspath(src))
    try:
        with open(src, encoding="utf-8") as f:
            c = json.load(f)
    except FileNotFoundError:
        print(f"[오류] 파일 없음: {src}"); sys.exit(2)
    except json.JSONDecodeError as ex:
        print(f"[오류] JSON 파싱 실패: {ex}"); sys.exit(2)
    if not isinstance(c, dict):
        print("[오류] 최상위가 객체(JSON object)가 아님"); sys.exit(2)

    errs = gate(c, preview=preview)
    if errs:
        print("=" * 60)
        print("렌더 거부 — 게이트 위반 (fidelity-guard 먼저 통과시키라)")
        print("=" * 60)
        for er in errs:
            print(f"  ✗ {er}")
        sys.exit(1)

    os.makedirs(out_dir, exist_ok=True)
    html_path = os.path.join(out_dir, "calendar.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(build_html(c, preview=preview))
    xlsx_path = build_xlsx(c, os.path.join(out_dir, "calendar.xlsx"))

    print("렌더 완료 ✓" + ("  (미리보기 — 미확정·최종 아님)" if preview else "  (목사 확정)"))
    print(f"  · {html_path}")
    print(f"  · {xlsx_path}" + ("  (openpyxl 없어 csv 폴백)" if xlsx_path.endswith(".csv") else ""))
    sys.exit(0)


if __name__ == "__main__":
    main()
